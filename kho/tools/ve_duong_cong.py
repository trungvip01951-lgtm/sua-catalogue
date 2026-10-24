#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ve_duong_cong.py — Vẽ biểu đồ Q–H và/hoặc bảng thông số bơm theo hồ sơ hãng.

Xuất PDF vector (chữ thật, phông nhúng dạng rút gọn), mỗi trang là một khối để mã dựng
catalogue chèn vào trang bằng page.show_pdf_page().

Căn cứ: claude/PHUONG_AN_DA_DUYET_20261013.md (Hồ sơ hãng, Chuẩn biểu đồ và bảng, Chuẩn chữ).
Mọi màu / phông / cỡ chữ / kiểu biểu đồ / kiểu bảng / khối điều kiện ĐỌC từ rules/HO_SO_HANG.json.

Số liệu: DU_LIEU/<NHÓM>.xlsx
  THONG_TIN: khoa | gia_tri | nguon      (dùng: so_cuc, toc_do_vg_ph)
  MODEL    : hang | model_1pha | model_3pha | p2_kw | p2_hp | dong_A_1pha | dong_A_3pha | ...
  QH       : hang | model | q_m3h | h_m | nguon   (model = model_3pha, nếu không có thì model_1pha)
  Cột hang: tên hãng; để trống / "*" / "CHUNG" = dùng cho cả 3 hãng; nhiều hãng cách nhau bằng dấu phẩy.

Dùng:
  python ve_duong_cong.py --du-lieu DU_LIEU/CM.xlsx --hang PEROTAC --ra cm.pdf
         [--model "CM 80-125/4,CM 80-125/5.5"] [--kieu bieu_do|bang|ca_hai]
         [--rong-pt 774 --cao-pt 527] [--ho-so rules/HO_SO_HANG.json] [--khong-tieu-de]
         [--cho-phep-phong-gan-giong] [--chap-nhan-gon] [--so-voi DU_LIEU/<NHÓM>.xlsx]

  Mặc định khổ = vùng chữ trang A4 của hãng (khổ trừ lề le_mm[0] hai bên).
  Kiểu "bang" không ghi --cao-pt: trang cao vừa đúng bảng.
  Số đường tối đa / biểu đồ: AQUARIS 7 (bảng màu Okabe–Ito bỏ vàng có 7 màu), PEROTAC/REDPUMP 8
  (so_duong_toi_da_moi_bieu_do). Nhiều hơn -> tự tách nhiều trang (mỗi trang một biểu đồ + bảng của nhóm).
  Phông tìm theo thứ tự: <thư mục script>/../fonts, <gốc kho>/vendor/fonts, hệ thống (đọc tên trong tệp).
  Thiếu phông đúng hồ sơ hãng -> dừng mã 3 "thiếu phông", trừ khi --cho-phep-phong-gan-giong.
  Dữ liệu tạo gợn (độ cong đổi chiều > 1 lần, hoặc đoạn gần phẳng rồi dốc gắt): so với số liệu gốc đã duyệt
  (--so-voi; không ghi thì DU_LIEU/<cùng tên tệp> nếu tệp vào nằm ngoài DU_LIEU/, còn tệp vào trong DU_LIEU/ thì
  so với chính nó). Dãy (Q,H) giống hệt bản gốc -> chỉ cảnh báo "gợn có sẵn trong số liệu đã duyệt — không chặn";
  điểm đã đổi (hoặc không có bản gốc) mà tạo gợn -> dừng mã 2, trừ khi --chap-nhan-gon.
  Nguồn không có dòng điện A ở model nào -> bỏ cột A (QĐ 152, ghi chú trong cảnh báo).
  Có A nhưng vài model thiếu -> giữ cột, ô trống + cảnh báo "THIẾU DÒNG ĐIỆN A" (điền số suy ra vào DU_LIEU theo QTH 4.6
  rồi vẽ lại; công cụ không tự suy, không điền 0).
  Mọi cảnh báo / lỗi ghi thêm vào <ra>.canh_bao.txt.

Mã thoát: 0 đạt · 2 dữ liệu sai (Q tăng H giảm không đạt, gợn, thiếu model/điểm) ·
          3 không vừa khổ / nhãn không đặt được / thiếu phông · 1 lỗi khác (hồ sơ, tham số).
"""
import argparse
import glob
import json
import math
import os
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pymupdf
from scipy.interpolate import PchipInterpolator, CubicSpline

LMIN = 1000.0 / 60.0          # 1 m³/h = 16,667 l/min
MM = 72.0 / 25.4
GOC = Path(__file__).resolve().parent.parent
HO_SO_MAC_DINH = GOC / 'rules' / 'HO_SO_HANG.json'
THU_MUC_DU_LIEU = GOC / 'DU_LIEU'          # số liệu gốc đã duyệt (dùng để so khi gặp gợn)
GIAN_BANG_TOI_DA = 3.0       # bảng giãn đều các cột tối đa 3 lần bề rộng tự nhiên
TRANG = (1.0, 1.0, 1.0)
DEN = (0.0, 0.0, 0.0)


class LoiVe(Exception):
    ma = 1


class LoiDuLieu(LoiVe):
    ma = 2


class LoiBoCuc(LoiVe):
    ma = 3


# ----------------------------------------------------------------------------- màu
def mau(h):
    h = str(h).strip().lstrip('#')
    return tuple(int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4))


def hexm(c):
    return '#%02X%02X%02X' % tuple(int(round(max(0.0, min(1.0, v)) * 255)) for v in c)


def _tt(v):
    return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4


def do_choi(c):
    r, g, b = (_tt(v) for v in c)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def tuong_phan(a, b):
    """Tỉ số tương phản WCAG 2.x giữa hai màu RGB (0..1)."""
    la, lb = do_choi(a), do_choi(b)
    if la < lb:
        la, lb = lb, la
    return (la + 0.05) / (lb + 0.05)


def tron(a, b, t):
    return tuple(x + (y - x) * t for x, y in zip(a, b))


def dam_den_dat(c, nen, nguong):
    """Làm đậm dần (giữ sắc) tới khi tương phản với nền >= ngưỡng."""
    t, c2 = 0.0, c
    while tuong_phan(c2, nen) < nguong and t < 1.0:
        t += 0.02
        c2 = tron(c, DEN, t)
    return c2


# ----------------------------------------------------------------------------- phông
THU_MUC_PHONG_HE_THONG = ('/usr/share/fonts', '/usr/local/share/fonts', '~/.fonts', '~/.local/share/fonts',
                          'C:/Windows/Fonts', '~/AppData/Local/Microsoft/Windows/Fonts', '/Library/Fonts',
                          '~/Library/Fonts')


def _chuan(s):
    return re.sub(r'[^a-z0-9]', '', str(s).lower())


def _ten_trong_tep(tep):
    try:
        return pymupdf.Font(fontfile=str(tep)).name
    except Exception:  # noqa: BLE001 — tệp hỏng / không phải phông
        return ''


def _dung_phong(tep, ho, dam):
    """Tệp có đúng họ + kiểu (Regular / Bold) không — đọc tên ghi TRONG tệp, không tin tên tệp."""
    muc = _chuan(ho) + ('bold' if dam else '')
    return _chuan(_ten_trong_tep(tep)) in ((muc,) if dam else (muc, muc + 'regular'))


def thu_muc_phong_rieng():
    """Thứ tự tìm: <thư mục chứa script>/../fonts, <gốc kho>/vendor/fonts (rồi mới tới hệ thống)."""
    return [Path(__file__).resolve().parent.parent / 'fonts', GOC / 'vendor' / 'fonts']


def tim_tep_phong(ho, dam, cho_phep_gan_giong=False):
    """Trả (tệp, tên phông thực, nguồn, gần_giống). Không có phông đúng họ: LoiBoCuc 'thiếu phông',
    trừ khi cho_phep_gan_giong (khi đó dùng phông máy đưa ra và đánh dấu gần giống)."""
    kieu = 'Bold' if dam else 'Regular'
    for d in thu_muc_phong_rieng():
        if d.is_dir():
            for p in sorted(d.iterdir()):
                if p.suffix.lower() in ('.ttf', '.otf') and _dung_phong(p, ho, dam):
                    return str(p), _ten_trong_tep(p), str(d), False
    ung_vien = []
    try:
        r = subprocess.run(['fc-match', '-f', '%{file}', f'{ho}:style={kieu}'],
                           capture_output=True, text=True, timeout=20)
        if r.returncode == 0 and r.stdout.strip():
            ung_vien.append(r.stdout.strip())
    except (OSError, subprocess.SubprocessError):
        pass
    for d in THU_MUC_PHONG_HE_THONG:
        d = os.path.expanduser(d)
        if os.path.isdir(d):
            ung_vien += sorted(glob.glob(f'{d}/**/*{ho.replace(" ", "")}*.[ot]tf', recursive=True))
    for p in ung_vien:
        if os.path.isfile(p) and p.lower().endswith(('.ttf', '.otf')) and _dung_phong(p, ho, dam):
            return p, _ten_trong_tep(p), 'hệ thống', False
    if not cho_phep_gan_giong:
        raise LoiBoCuc(f'thiếu phông {ho} {kieu} — không có trong fonts/, vendor/fonts/ hay phông hệ thống. '
                       f'Chép phông vào vendor/fonts/ hoặc chạy lại với --cho-phep-phong-gan-giong.')
    thay = [p for p in ung_vien if os.path.isfile(p) and p.lower().endswith(('.ttf', '.otf'))]
    for d in thu_muc_phong_rieng():
        if d.is_dir():
            thay += [str(p) for p in sorted(d.iterdir()) if p.suffix.lower() in ('.ttf', '.otf')]
    for d in THU_MUC_PHONG_HE_THONG:
        d = os.path.expanduser(d)
        if os.path.isdir(d):
            thay += sorted(glob.glob(f'{d}/**/*.ttf', recursive=True))
    for p in thay:
        ten = _ten_trong_tep(p)
        if ten and ('bold' in ten.lower()) == dam and 'italic' not in ten.lower() and 'oblique' not in ten.lower():
            return p, ten, 'thay thế', True
    raise LoiBoCuc(f'thiếu phông {ho} {kieu} và không tìm được phông thay thế nào trên máy.')


class Phong:
    """Một họ phông (thường + đậm). Tìm ở fonts/, vendor/fonts/, rồi hệ thống; thiếu thì báo (mã 3)."""

    def __init__(self, ho, cho_phep_gan_giong=False):
        self.ho = ho
        self.tep, self.ho_thuc, self.nguon, self.font, self.asc, self.desc = {}, {}, {}, {}, {}, {}
        self.gan_giong = False
        for dam in (False, True):
            tep, ten, nguon, gan = tim_tep_phong(ho, dam, cho_phep_gan_giong)
            self.tep[dam], self.ho_thuc[dam], self.nguon[dam] = tep, ten, nguon
            self.gan_giong |= gan
            f = pymupdf.Font(fontfile=tep)
            self.font[dam], self.asc[dam], self.desc[dam] = f, f.ascender, f.descender
        self.ten = {False: 'VeR', True: 'VeB'}

    def co_glyph(self, ch, dam=False):
        return ch.isspace() or self.font[dam].has_glyph(ord(ch)) != 0

    def rong(self, s, co, dam=False):
        return self.font[dam].text_length(s, fontsize=co)

    def cao(self, co, dam=False):
        return (self.asc[dam] - self.desc[dam]) * co

    def chan(self, tam_y, co, dam=False):
        """Đường chân chữ để khung chữ có tâm dọc = tam_y."""
        return tam_y + (self.asc[dam] + self.desc[dam]) / 2.0 * co

    def mo_ta(self):
        s = (f'{self.ho_thuc[False]} + {self.ho_thuc[True]} ({os.path.basename(self.tep[False])}, '
             f'{os.path.basename(self.tep[True])}; nguồn: {self.nguon[False]})')
        if self.gan_giong:
            return f'GẦN GIỐNG: {s} thay cho "{self.ho}" (hồ sơ hãng)'
        return s + ' — đúng hồ sơ hãng'


SUP = {'⁰': '0', '¹': '1', '²': '2', '³': '3', '⁴': '4', '⁵': '5', '⁶': '6', '⁷': '7',
       '⁸': '8', '⁹': '9', '⁻': '−', '⁺': '+'}


# ----------------------------------------------------------------------------- hồ sơ hãng
BIEN_THE_NHAN = {
    # kiểu nhãn (json) -> bố trí phụ suy ra: thang l/min, khung trục
    'so_khoanh_tron_va_chu_giai': {'lmin': 'tren', 'khung': 'trai_duoi'},
    'ten_model_dau_trai_duong': {'lmin': 'duoi', 'khung': 'day_du'},
    'ten_model_trong_o_mau_cuoi_phai_duong': {'lmin': 'tren', 'khung': 'duoi'},
}
KIEU_MAU = ('moi_model_mot_mau_okabe_ito', 'mot_tong_xanh', 'do_xam')
KIEU_BANG = ('ngang_manh', 'day_du_manh', 'dau_bang_toi_chu_trang_soc_xam')
KIEU_HUONG = ('model_theo_cot', 'model_theo_dong')
KIEU_DO_HOA = ('dai_co_y', 'bang_dau_trang_xanh_nhat', 'cot_toi_trai_28_30_phan_tram')


class KieuHang:
    def __init__(self, HS, hang):
        ch = HS['chung']
        h = HS['hang'][hang]
        self.hang = hang
        co = ch['co_chu_pt']
        self.co_so = float(max(co['bang_nen'], co['bang_min'], co['nhan_truc_min'], co['nhan_duong_min']))
        self.co_min_bang = float(co['bang_min'])
        self.co_min_nhan = float(max(co['nhan_truc_min'], co['nhan_duong_min'], co['bang_min']))
        self.co_than = float(co['than_min'])
        self.co_khoi = self.co_than + 1.0          # tiêu đề khối (không phải tiêu đề trang)
        self.net = max(float(ch['net_duong_pt_min']), 1.5) + 0.25
        self.luoi = float(ch['luoi_chinh_pt'])
        self.toi_da = int(ch['so_duong_toi_da_moi_bieu_do'])
        self.tp_chu = float(ch['tuong_phan']['chu'])
        self.tp_net = float(ch['tuong_phan']['chu_lon_va_net'])
        self.dieu_kien = ch['khoi_dieu_kien_duong_cong']
        self.dien = ch['dien']
        self.okabe = [mau(x) for x in ch['okabe_ito_bo_vang']]
        self.chu_cam = [s.lower() for s in ch.get('chu_cam', [])]
        kho = ch['kho_' + h['kho']]
        le = float(ch['le_mm'][0]) * MM
        self.rong_md, self.cao_md = round(kho[0] - 2 * le, 1), round(kho[1] - 2 * le, 1)
        self.phong = h['phong'][0]
        self.m = {k: mau(v) for k, v in h['mau'].items() if isinstance(v, str)}
        self.co_y = [mau(x) for x in h['mau'].get('co_y', [])]
        bd = h['bieu_do']
        self.kieu_mau, self.kieu_nhan, self.truc_h, self.nen_bd = bd['mau_duong'], bd['nhan'], bd['truc_H'], mau(bd['nen'])
        self.bang_huong, self.bang_ke = h['bang']['huong'], h['bang']['ke']
        self.do_hoa = h['do_hoa_phu']
        for gt, ds, ten in ((self.kieu_nhan, BIEN_THE_NHAN, 'bieu_do.nhan'), (self.kieu_mau, KIEU_MAU, 'bieu_do.mau_duong'),
                            (self.bang_ke, KIEU_BANG, 'bang.ke'), (self.bang_huong, KIEU_HUONG, 'bang.huong'),
                            (self.do_hoa, KIEU_DO_HOA, 'do_hoa_phu'), (self.truc_h, ('trai', 'phai'), 'bieu_do.truc_H')):
            if gt not in ds:
                raise LoiVe(f'{hang}: kiểu "{gt}" ở {ten} chưa được ve_duong_cong.py hỗ trợ.')
        self.bt = BIEN_THE_NHAN[self.kieu_nhan]
        self.chu = self.m['chinh']
        self.luoi_mau = tron(self.nen_bd, self.m.get('duong', self.chu), 0.18)
        self.toi_da_bd = min(self.toi_da, len(self.okabe)) if self.kieu_mau == 'moi_model_mot_mau_okabe_ito' else self.toi_da


def mau_duong(S, n, canh_bao):
    """Danh sách màu đường cho n model, theo kiểu hãng; nét đạt >= 3:1 trên nền biểu đồ."""
    if S.kieu_mau == 'moi_model_mot_mau_okabe_ito':
        if n > len(S.okabe):
            raise LoiBoCuc(f'{n} model > {len(S.okabe)} màu Okabe–Ito — cần tách biểu đồ.')
        goc = S.okabe[:n]
        nguong = S.tp_net
    elif S.kieu_mau == 'mot_tong_xanh':
        c0, c1 = S.m['chinh'], S.m.get('duong', S.m['chinh'])
        day = [tron(c0, c1, t) for t in np.linspace(0, 1, 6)] + [tron(c1, TRANG, t) for t in (0.1, 0.2)]
        chon = [day[int(round(k))] for k in np.linspace(0, len(day) - 1, n)] if n > 1 else [c1]
        thu_tu = list(range(0, n, 2)) + list(range(1, n, 2))   # xen kẽ đậm/nhạt cho 2 đường kề nhau
        goc = [chon[k] for k in thu_tu]
        nguong = S.tp_chu                                       # nhãn chữ cùng màu đường
    else:  # do_xam
        do, xam = S.m['nhan'], S.m['chinh']
        day = [do, tron(xam, TRANG, 0.3), tron(do, DEN, 0.35), xam,
               tron(do, TRANG, 0.15), tron(xam, TRANG, 0.15), tron(do, DEN, 0.55), tron(xam, DEN, 0.4)]
        goc = [day[k % len(day)] for k in range(n)]
        nguong = S.tp_chu                                       # chữ trắng trong ô màu
    ra = []
    for c in goc:
        c2 = dam_den_dat(c, S.nen_bd, nguong)
        if S.kieu_mau == 'do_xam':
            c2 = dam_den_dat(c2, TRANG, nguong)
        if hexm(c2) != hexm(c):
            canh_bao.append(f'Màu {hexm(c)} làm đậm thành {hexm(c2)} để đạt tương phản {nguong:g}:1.')
        ra.append(c2)
    return ra


# ----------------------------------------------------------------------------- số liệu
def _so(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '.')
    if s in ('', '–', '-', '—'):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _chu(v):
    return '' if v is None else str(v).strip()


def chuan_ten(s):
    return re.sub(r'\s+', ' ', _chu(s)).upper()


def fmt(v):
    return '' if v is None else ('%.6g' % v)


def khop_hang(o, hang):
    s = _chu(o).upper()
    if s in ('', '*', 'CHUNG', 'ALL', 'CA 3', 'CẢ 3', 'CA_3'):
        return True
    return hang in [t.strip() for t in re.split(r'[,;/|+]', s)]


def _sheet(wb, ten):
    if ten not in wb.sheetnames:
        raise LoiDuLieu(f'Tệp số liệu thiếu sheet "{ten}".')
    rows = list(wb[ten].iter_rows(values_only=True))
    if not rows:
        return []
    hdr = [_chu(c).lower() for c in rows[0]]
    return [(i, dict(zip(hdr, r))) for i, r in enumerate(rows[1:], start=2)
            if any(c not in (None, '') for c in r)]


def doc_du_lieu(duong_dan, hang):
    p = Path(duong_dan)
    if not p.is_file():
        raise LoiDuLieu(f'Không thấy tệp số liệu: {p}')
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    tt = {}
    for _, r in _sheet(wb, 'THONG_TIN'):
        k = _chu(r.get('khoa'))
        if k:
            tt[k] = r.get('gia_tri')
    ds, theo_ten = [], {}
    for dong, r in _sheet(wb, 'MODEL'):
        if not khop_hang(r.get('hang'), hang):
            continue
        m1, m3 = _chu(r.get('model_1pha')), _chu(r.get('model_3pha'))
        if not (m1 or m3):
            continue
        md = dict(khoa=m3 or m1, m1=m1, m3=m3, kw=_so(r.get('p2_kw')), hp=_so(r.get('p2_hp')),
                  a1=_so(r.get('dong_a_1pha')), a3=_so(r.get('dong_a_3pha')), dong=dong, diem=[])
        for t in {m1, m3} - {''}:
            k = chuan_ten(t)
            if k in theo_ten:
                raise LoiDuLieu(f'MODEL dòng {dong}: model "{t}" của {hang} bị trùng (đã có ở dòng {theo_ten[k]["dong"]}).')
            theo_ten[k] = md
        ds.append(md)
    for dong, r in _sheet(wb, 'QH'):
        if not khop_hang(r.get('hang'), hang):
            continue
        md = theo_ten.get(chuan_ten(r.get('model')))
        if md is None:
            continue
        q, h = _so(r.get('q_m3h')), _so(r.get('h_m'))
        if q is None or h is None:
            continue
        md['diem'].append((q, h, dong))
    wb.close()
    return tt, ds, theo_ten


def chon_model(ds, theo_ten, ten_chon, hang):
    if not ten_chon:
        if not ds:
            raise LoiDuLieu(f'Tệp số liệu không có model nào cho hãng {hang}.')
        return ds
    out = []
    for t in ten_chon:
        md = theo_ten.get(chuan_ten(t))
        if md is None:
            raise LoiDuLieu(f'Không có model "{t}" cho hãng {hang} trong sheet MODEL.')
        if md not in out:
            out.append(md)
    return out


def kiem_duong(md, can_duong):
    """Q tăng thì H không tăng; không trùng Q; qua đúng điểm, không móc. Trả về hàm nội suy:
    spline bậc 3 mềm (QĐ 158) nếu giữ được chiều giảm và không vượt điểm, nếu không thì PCHIP."""
    ten = md['khoa']
    pts = sorted(md['diem'], key=lambda p: p[0])
    if not pts:
        raise LoiDuLieu(f'Model {ten}: không có điểm Q–H nào trong sheet QH.')
    for q, h, dong in pts:
        if q < 0 or h < 0:
            raise LoiDuLieu(f'Model {ten}: điểm âm (Q {q:g} m³/h; H {h:g} m) ở QH dòng {dong}.')
    for (q1, h1, d1), (q2, h2, d2) in zip(pts, pts[1:]):
        if q2 == q1:
            raise LoiDuLieu(f'Model {ten}: trùng Q = {q1:g} m³/h (H {h1:g} và {h2:g} m; QH dòng {d1}, {d2}).')
        if h2 > h1:
            raise LoiDuLieu(f'Model {ten}: H TĂNG khi Q tăng — điểm (Q {q1:g} m³/h; H {h1:g} m) → '
                            f'(Q {q2:g} m³/h; H {h2:g} m), QH dòng {d1} → {d2}. Dừng, không vẽ.')
    md['diem'] = pts
    if not can_duong:
        return None
    if len(pts) < 2:
        raise LoiDuLieu(f'Model {ten}: chỉ có {len(pts)} điểm Q–H, cần ≥ 2 để vẽ đường.')
    q = np.array([p[0] for p in pts])
    h = np.array([p[1] for p in pts])
    f = PchipInterpolator(q, h)
    qq = np.linspace(q[0], q[-1], 2000)
    hh = f(qq)
    if np.any(np.diff(hh) > 1e-9) or hh.max() > h.max() + 1e-9 or hh.min() < h.min() - 1e-9:
        raise LoiDuLieu(f'Model {ten}: đường nội suy có móc/vượt điểm — kiểm lại số liệu.')
    if np.max(np.abs(f(q) - h)) > 1e-6:
        raise LoiDuLieu(f'Model {ten}: đường không qua đúng điểm bảng.')
    # QĐ 158 (23/10/2026) "đường đặc tính vẽ cong mềm mại": ưu tiên spline bậc 3 (not-a-knot) qua đúng điểm —
    # độ cong liền mạch, không bẻ góc tại điểm làm tròn. Chỉ nhận khi Q tăng H không tăng và không vượt điểm;
    # không đạt thì giữ PCHIP (giữ chiều, không móc).
    # Thử lần lượt: not-a-knot; rồi đầu đường dốc 0 tại Q đầu (dáng bơm ly tâm ở điểm đóng van).
    # Chỉ nhận spline khi nó chỉ làm mềm góc, không đổi dáng (mục 4.2): không tăng, không vượt điểm,
    # số lần độ cong đổi chiều ≤ PCHIP và ≤ 1, lệch PCHIP ≤ 2 % biên độ H. Không đạt thì giữ PCHIP.
    if len(pts) >= 4:
        bien = max(h.max() - h.min(), 1e-9)
        def doi_chieu(g):
            c = g(qq, 2); c = c[np.abs(c) > 1e-3 * bien / max(q[-1] - q[0], 1e-9) ** 2]
            return int(np.sum(np.diff(np.sign(c)) != 0))
        dc_p = doi_chieu(f)
        for bc in ('not-a-knot', ((1, 0.0), 'not-a-knot')):
            s = CubicSpline(q, h, bc_type=bc)
            hs = s(qq)
            if np.any(np.diff(hs) > 1e-9) or hs.max() > h.max() + 1e-9 or hs.min() < h.min() - 1e-9:
                continue
            if np.max(np.abs(hs - hh)) > 0.02 * bien or doi_chieu(s) > min(dc_p, 1):
                continue
            return s
    return f


GON_LAM_TRON = 1.0     # điểm giữa lệch khỏi dây cung 2 điểm kề phải > 1 bước làm tròn của dãy H (1 / 0,5 / 0,1 m)
GON_TI_LE_H = 0.02     # ... và > 2 % biên độ H của model (≈ 5 pt trên vùng vẽ cao 250 pt) mới tính là gợn thấy được


def _buoc_lam_tron(h):
    if np.all(np.abs(h - np.round(h)) < 1e-9):
        return 1.0
    if np.all(np.abs(h * 2 - np.round(h * 2)) < 1e-9):
        return 0.5
    return 0.1


def kiem_gon(md):
    """Dữ liệu tạo gợn (đường PCHIP lượn sóng): (a) độ cong (sai phân bậc 2) đổi chiều quá 1 lần;
    (b) đoạn gần phẳng giữa dãy rồi dốc gắt. Chỉ tính chỗ lệch thấy được (GON_LAM_TRON, GON_TI_LE_H).
    Trả danh sách mô tả (rỗng = không gợn)."""
    pts = md['diem']
    if len(pts) < 4:
        return []
    q = np.array([p[0] for p in pts], float)
    h = np.array([p[1] for p in pts], float)
    dq = np.diff(q)
    m = -np.diff(h) / dq                                         # độ dốc xuống của từng đoạn
    tb = (h[0] - h[-1]) / (q[-1] - q[0])
    nguong = max(GON_LAM_TRON * _buoc_lam_tron(h), GON_TI_LE_H * (h.max() - h.min()))
    de = [(m[i + 1] - m[i]) * dq[i] * dq[i + 1] / (dq[i] + dq[i + 1]) for i in range(len(m) - 1)]  # ứng điểm i+1

    def p(k):
        return f'(Q {q[k]:g} m³/h; H {h[k]:g} m)'

    loi = []
    for i in range(1, len(m) - 1):
        if m[i] < 0.35 * max(m[i - 1], tb) and m[i + 1] > 4 * max(m[i], 1e-9) and -de[i - 1] > nguong and de[i] > nguong:
            loi.append(f'đoạn gần phẳng rồi dốc gắt {p(i)} → {p(i + 1)} → {p(i + 2)}')
    dau = [(np.sign(d), i + 1) for i, d in enumerate(de) if abs(d) > nguong]
    doi = [b[1] for a, b in zip(dau, dau[1:]) if a[0] != b[0]]
    if len(doi) > 1:
        loi.append(f'độ cong đổi chiều {len(doi)} lần tại ' + ', '.join(p(k) for k in doi))
    return loi


def chia_nhom(ds, toi_da):
    n = len(ds)
    k = max(1, math.ceil(n / toi_da))
    out, i = [], 0
    for j in range(k):
        sz = n // k + (1 if j < n % k else 0)
        out.append(ds[i:i + sz])
        i += sz
    return out


def _gia_tri_hang(v, hang, ten, canh_bao):
    """Số từ THONG_TIN: số trơn, hoặc chuỗi kiểu "AQUARIS: 2850/2900 vg/ph; PEROTAC: 2900 vg/ph"."""
    if v is None or isinstance(v, (int, float)):
        return _so(v)
    s = str(v)
    phan = [x for x in re.split(r'[;\n]', s) if x.strip().upper().startswith(hang)]
    so = re.findall(r'\d+(?:[.,]\d+)?', (phan[0].split(':', 1)[-1] if phan else s))
    if len(so) == 1:
        return float(so[0].replace(',', '.'))
    if so:
        canh_bao.append(f'THONG_TIN {ten} của {hang} có nhiều giá trị ("{(phan[0] if phan else s).strip()}") '
                        f'— khối điều kiện giữ số mặc định của hồ sơ hãng.')
    return None


def khoi_dieu_kien(S, tt, canh_bao):
    s = S.dieu_kien
    cuc = _gia_tri_hang(tt.get('so_cuc'), S.hang, 'so_cuc', canh_bao)
    n = _gia_tri_hang(tt.get('toc_do_vg_ph'), S.hang, 'toc_do_vg_ph', canh_bao)
    if cuc:
        s = re.sub(r'\b\d+\s*poles\b', f'{cuc:g} poles', s)
    if n:
        s = re.sub(r'≈\s*\d+(?:[.,]\d+)?\s*min', f'≈ {n:g} min', s)
    return s


# ----------------------------------------------------------------------------- trục
def cac_buoc(vmax):
    e = math.floor(math.log10(max(vmax, 1e-9))) - 3
    return [m * 10.0 ** k for k in range(e, e + 7) for m in (1, 2, 2.5, 5)]


def chon_truc(vmax, do_dai, rong_nhan, kc_min, toi_da_vach=14):
    """Bước 'đẹp' nhỏ nhất để 2 nhãn cạnh nhau cách >= kc_min và không chạm nhau. Trả (bước, đỉnh)."""
    for b in cac_buoc(vmax):
        dinh = math.ceil(vmax / b - 1e-9) * b
        n = int(round(dinh / b))
        if n > toi_da_vach or n < 1:
            continue
        kc = do_dai * b / dinh
        if kc >= kc_min and kc >= rong_nhan([i * b for i in range(n + 1)]) + 6:
            return b, dinh
    raise LoiBoCuc('Trục quá ngắn để ghi số — cần tăng kích thước biểu đồ.')


def chon_buoc_phu(vmax, pt_moi_don_vi, rong_nhan, kc_min, toi_da_vach=14):
    for b in cac_buoc(vmax):
        n = int(math.floor(vmax / b + 1e-9))
        if n > toi_da_vach or n < 1:
            continue
        kc = b * pt_moi_don_vi
        if kc >= kc_min and kc >= rong_nhan([i * b for i in range(n + 1)]) + 6:
            return b
    raise LoiBoCuc('Thang l/min quá ngắn để ghi số.')


# ----------------------------------------------------------------------------- khung vẽ
class Khung:
    """Một trang PDF đang vẽ + sổ ghi mọi chữ (khung, tương phản) để tự kiểm trước khi lưu."""

    def __init__(self, doc, w, h, ph, S, dung=None):
        self.page = doc.new_page(width=w, height=h)
        self.ph, self.S = ph, S
        self.dung = dung if dung is not None else {False: set(), True: set()}   # ký tự đã dùng / độ đậm
        for dam in (False, True):
            self.page.insert_font(fontname=ph.ten[dam], fontfile=ph.tep[dam])
        self.ghi, self.loi = [], []

    def _chia(self, s, co, dam):
        ph, out, buf, i, n = self.ph, [], '', 0, len(s)
        while i < n:
            ch = s[i]
            if ch in SUP:
                j = i
                while j < n and s[j] in SUP:
                    j += 1
                seg = s[i:j]
                if all(ph.co_glyph(c, dam) for c in seg):
                    buf += seg
                else:                                   # phông thiếu chữ số mũ -> chữ nhỏ nâng lên
                    if buf:
                        out.append((buf, co, 0.0, ph.rong(buf, co, dam)))
                        buf = ''
                    sub = ''.join(SUP[c] for c in seg)
                    sub = ''.join(c if ph.co_glyph(c, dam) else '-' for c in sub)
                    cs = min(co, max(self.S.co_min_nhan, round(co * 0.72, 2)))
                    out.append((sub, cs, co * 0.36, ph.rong(sub, cs, dam)))
                i = j
                continue
            if not ph.co_glyph(ch, dam):
                raise LoiVe(f'Phông "{ph.ho}" thiếu ký tự "{ch}" (U+{ord(ch):04X}) trong "{s}".')
            buf += ch
            i += 1
        if buf:
            out.append((buf, co, 0.0, ph.rong(buf, co, dam)))
        return out

    def rong(self, s, co, dam=False):
        return sum(r[3] for r in self._chia(s, co, dam))

    def chu(self, x, y, s, co, mau_chu, nen, dam=False, canh='trai', xoay=0, loai='chu'):
        """Viết chữ. xoay=0: (x, y) = điểm canh trên đường chân chữ. xoay=90: (x, y) = tâm khung chữ,
        chữ đọc từ dưới lên. Trả về khung chữ (Rect)."""
        if not s:
            return None
        ph = self.ph
        runs = self._chia(s, co, dam)
        for t, *_ in runs:
            self.dung[dam].update(t)
        w = sum(r[3] for r in runs)
        asc, desc, ten = ph.asc[dam], ph.desc[dam], ph.ten[dam]
        if xoay == 0:
            x0 = x - w if canh == 'phai' else (x - w / 2.0 if canh == 'giua' else x)
            cx, top, bot = x0, y - asc * co, y - desc * co
            for t, c, rise, rw in runs:
                self.page.insert_text((cx, y - rise), t, fontsize=c, fontname=ten, color=mau_chu)
                top, bot = min(top, y - rise - asc * c), max(bot, y - rise - desc * c)
                cx += rw
            r = pymupdf.Rect(x0, top, x0 + w, bot)
        else:
            bx = x + (asc + desc) / 2.0 * co
            by = y + w / 2.0
            cy = by
            for t, c, rise, rw in runs:
                self.page.insert_text((bx - rise, cy), t, fontsize=c, fontname=ten, color=mau_chu, rotate=90)
                cy -= rw
            r = pymupdf.Rect(bx - asc * co, by - w, bx - desc * co, by)
        tp = tuong_phan(mau_chu, nen)
        self.ghi.append(dict(t=s, r=r, tp=tp, co=co, loai=loai))
        if tp < self.S.tp_chu - 1e-9:
            self.loi.append(f'Tương phản {tp:.2f}:1 < {self.S.tp_chu:g}:1 — chữ "{s}" {hexm(mau_chu)} trên nền {hexm(nen)}')
        for c in self.S.chu_cam:
            if re.search(r'(?<![a-z])' + re.escape(c) + r'(?![a-z])', s.lower()):
                self.loi.append(f'Chữ cấm "{c}" trong "{s}"')
        return r

    def chu_giua_o(self, r, s, co, mau_chu, nen, dam=False, canh='giua', dem=4.0, loai='bang'):
        y = self.ph.chan((r.y0 + r.y1) / 2.0, co, dam)
        if canh == 'phai':
            return self.chu(r.x1 - dem, y, s, co, mau_chu, nen, dam, 'phai', loai=loai)
        if canh == 'trai':
            return self.chu(r.x0 + dem, y, s, co, mau_chu, nen, dam, 'trai', loai=loai)
        return self.chu((r.x0 + r.x1) / 2.0, y, s, co, mau_chu, nen, dam, 'giua', loai=loai)

    def tu_kiem(self):
        loi = list(self.loi)
        g = self.ghi
        for i in range(len(g)):
            for j in range(i + 1, len(g)):
                a, b = g[i]['r'], g[j]['r']
                ox = min(a.x1, b.x1) - max(a.x0, b.x0)
                oy = min(a.y1, b.y1) - max(a.y0, b.y0)
                if ox > 0.05 and oy > 0.05:
                    loi.append(f'Chữ chồng nhau: "{g[i]["t"]}" và "{g[j]["t"]}"')
        if loi:
            raise LoiBoCuc('Tự kiểm không đạt:\n  - ' + '\n  - '.join(loi[:20]))
        return min((x['tp'] for x in g), default=21.0)


def sua_tounicode(doc, ph, dung):
    """Ghi lại bảng ToUnicode của phông nhúng: mỗi glyph -> đúng ký tự đã viết. MuPDF lấy ánh xạ ngược của cả phông,
    với Lato ra dấu cách = U+00A0, '-' = U+2010, '·' = U+2219 -> tìm/chép chữ trong PDF sai (vd. mã model)."""
    da_sua = set()
    for page in doc:
        for f in page.get_fonts(full=True):
            xref, ten = f[0], f[4]
            dam = {ph.ten[False]: False, ph.ten[True]: True}.get(ten)
            if dam is None or xref in da_sua:
                continue
            m = re.search(r'/ToUnicode (\d+) 0 R', doc.xref_object(xref))
            if not m:
                continue
            cap = {}
            for ch in sorted(dung[dam]):
                gid = ph.font[dam].has_glyph(ord(ch))
                if gid and ord(ch) <= 0xFFFF and gid not in cap:
                    cap[gid] = ch
            dong = []
            ds = sorted(cap.items())
            for i in range(0, len(ds), 100):
                khoi = ds[i:i + 100]
                dong.append(f'{len(khoi)} beginbfchar')
                dong += [f'<{g:04X}> <{ord(c):04X}>' for g, c in khoi]
                dong.append('endbfchar')
            cmap = ('/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n'
                    '/CIDSystemInfo <</Registry(Adobe)/Ordering(UCS)/Supplement 0>> def\n'
                    '/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n'
                    '1 begincodespacerange\n<0000> <FFFF>\nendcodespacerange\n' + '\n'.join(dong) +
                    '\nendcmap\nCMapName currentdict /CMap defineresource pop\nend\nend\n')
            doc.update_stream(int(m.group(1)), cmap.encode('ascii'))
            da_sua.add(xref)


# ----------------------------------------------------------------------------- tiêu đề khối
def cao_tieu_de(S, ph):
    hc = ph.cao(S.co_khoi, True)
    return hc + (8 if S.do_hoa in ('bang_dau_trang_xanh_nhat', 'dai_co_y') else 6)


def ve_tieu_de(K, S, x0, y0, w, trai, phai=''):
    ph, co = K.ph, S.co_khoi
    h = cao_tieu_de(S, ph)
    if S.do_hoa == 'bang_dau_trang_xanh_nhat':
        nen = S.m.get('nen', TRANG)
        K.page.draw_rect(pymupdf.Rect(x0, y0, x0 + w, y0 + h), color=None, fill=nen, width=0)
        tam, xt = y0 + h / 2.0, x0 + 6
    elif S.do_hoa == 'cot_toi_trai_28_30_phan_tram':
        nen = TRANG
        K.page.draw_rect(pymupdf.Rect(x0, y0 + 1, x0 + 4, y0 + h - 1), color=None, fill=S.m['nhan'], width=0)
        tam, xt = y0 + h / 2.0, x0 + 10
    else:  # dai_co_y
        nen = TRANG
        tam, xt = y0 + (h - 5) / 2.0, x0
        for k, c in enumerate(S.co_y):
            K.page.draw_rect(pymupdf.Rect(x0 + 12 * k, y0 + h - 3, x0 + 12 * (k + 1), y0 + h - 0.8),
                             color=None, fill=c, width=0)
        if S.co_y and do_choi(S.co_y[1 % len(S.co_y)]) > 0.9:        # dải trắng của cờ: viền mảnh
            K.page.draw_rect(pymupdf.Rect(x0, y0 + h - 3, x0 + 12 * len(S.co_y), y0 + h - 0.8),
                             color=tron(S.chu, TRANG, 0.6), width=0.25)
    K.chu(xt, ph.chan(tam, co, True), trai, co, S.chu, nen, dam=True, loai='tieu_de')
    if phai:
        xr = x0 + w - (6 if S.do_hoa == 'bang_dau_trang_xanh_nhat' else 0)
        K.chu(xr, ph.chan(tam, S.co_so), phai, S.co_so, S.chu, nen, canh='phai', loai='tieu_de')
    return h


# ----------------------------------------------------------------------------- nhãn đường
class ODat:
    """Kiểm một khung nhãn có trống không: trong vùng vẽ, không chạm đường nào, không chạm nhãn khác."""

    def __init__(self, P, dem_duong):
        self.P, self.dem_d = P, dem_duong
        self.duong, self.nhan = [], []

    def tu_do(self, r, bo_qua=(), dem_n=1.5, dem_d=None):
        P = self.P
        if r.x0 < P.x0 + 1 or r.x1 > P.x1 - 1 or r.y0 < P.y0 + 1 or r.y1 > P.y1 - 1:
            return False
        d = self.dem_d if dem_d is None else dem_d
        for i, A in self.duong:
            if i in bo_qua:
                continue
            if np.any((A[:, 0] > r.x0 - d) & (A[:, 0] < r.x1 + d) & (A[:, 1] > r.y0 - d) & (A[:, 1] < r.y1 + d)):
                return False
        for q in self.nhan:
            if q.x0 < r.x1 + dem_n and r.x0 < q.x1 + dem_n and q.y0 < r.y1 + dem_n and r.y0 < q.y1 + dem_n:
                return False
        return True


def _doan(A, x0, x1):
    xs = A[:, 0]
    seg = A[(xs >= x0 - 0.5) & (xs <= x1 + 0.5)]
    if len(seg) == 0:
        seg = A[[int(np.argmin(np.abs(xs - (x0 + x1) / 2.0)))]]
    return seg[:, 1].min(), seg[:, 1].max()


def dat_trai(O, A, w, h, gap):
    """PEROTAC: tên model ở đầu trái đường (trên, rồi dưới), trượt dần sang phải nếu vướng."""
    xs = A[:, 0]
    x_s, x_e = xs[0], xs[-1]
    for d in [2, 5, 9, 14, 20, 28, 38, 50, 64, 80] + list(range(96, int(x_e - x_s) + 1, 12)):
        x0 = x_s + d
        if x0 > x_s + 0.6 * (x_e - x_s) and x0 + w > x_e:
            break
        lo, hi = _doan(A, x0, x0 + w)
        for r in (pymupdf.Rect(x0, lo - gap - h, x0 + w, lo - gap), pymupdf.Rect(x0, hi + gap, x0 + w, hi + gap + h)):
            if O.tu_do(r):
                return r
    return None


def dat_phai_gan(O, A, w, h, gap):
    """REDPUMP, khi bên phải điểm cuối bị vướng: đặt ô trên/dưới đoạn CUỐI đường, chỉ lùi trái tối đa
    max(0,6 bề rộng ô; 10 % bề ngang đường) để ô vẫn ở cuối đường."""
    xs = A[:, 0]
    x_e = A[-1, 0]
    gioi_han = max(0.6 * w, 0.1 * (x_e - xs[0]))
    for d in (0, 3, 6, 10, 15, 21, 28, 36, 45, 56, 70):
        if d > gioi_han:
            break
        x1 = x_e - d
        x0 = x1 - w
        if x0 < xs[0]:
            break
        lo, hi = _doan(A, x0, x1)
        for r in (pymupdf.Rect(x0, lo - gap - h, x1, lo - gap), pymupdf.Rect(x0, hi + gap, x1, hi + gap + h)):
            if O.tu_do(r):
                return r
    return None


def _dan_doc(ys, h, sep, y_min, y_max):
    """Dàn dọc các ô cùng cột: giữ thứ tự, không chồng, mỗi cụm đặt quanh trung bình vị trí lý tưởng."""
    s = h + sep
    cum = []                                     # [chỉ số đầu, số ô, tâm ô đầu]
    for k, y in enumerate(ys):
        cum.append([k, 1, y])
        while len(cum) > 1 and cum[-2][2] + cum[-2][1] * s > cum[-1][2] + 1e-9:
            k0, nn = cum[-2][0], cum[-2][1] + cum[-1][1]
            cum[-2:] = [[k0, nn, sum(ys[k0 + j] - j * s for j in range(nn)) / nn]]
    out = []
    for k0, nn, t in cum:
        t = max(y_min + h / 2, min(t, y_max - h / 2 - (nn - 1) * s))
        out += [t + j * s for j in range(nn)]
    return out


def _dan_qua_nhan(O, a, b):
    for q in O.nhan:
        goc = np.array([[q.x0, q.y0], [q.x1, q.y0], [q.x0, q.y1], [q.x1, q.y1], [(q.x0 + q.x1) / 2, (q.y0 + q.y1) / 2]])
        if _cat_doan(goc, a, b, 1.0):
            return True
        for t in np.linspace(0, 1, 25):
            x, y = a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t
            if q.x0 - 0.5 <= x <= q.x1 + 0.5 and q.y0 - 0.5 <= y <= q.y1 + 0.5:
                return True
    return False


def _thu_o_phai(O, i, cuoi, w, h, gap, c, ep_dan=False):
    """Ô tên bên phải điểm cuối, tâm dọc c. Lệch > nửa ô (hoặc cả cột có ô lệch): lùi thêm 6 pt và kéo
    đường dẫn ngắn từ điểm cuối."""
    x_e, y_e = cuoi
    lech = ep_dan or abs(c - y_e) > 0.5 * h
    x0 = x_e + gap + (6.0 if lech else 0.0)
    r = pymupdf.Rect(x0, c - h / 2, x0 + w, c + h / 2)
    if not O.tu_do(r):
        return None
    if not lech:
        return r, None
    dan = (x_e, y_e, r.x0 - 0.8, min(max(y_e, r.y0 + 2), r.y1 - 2))
    khac = [B for k, B in O.duong if k != i]
    if khac:
        Pk = np.vstack(khac)
        Ek = np.array([B[j] for B in khac for j in (0, -1)])
        if _cat_doan(Pk, dan[:2], dan[2:], 2.0) or _cat_doan(Ek, dan[:2], dan[2:], 3.0):
            return None
    if _dan_qua_nhan(O, dan[:2], dan[2:]):
        return None
    return r, dan


def dat_cuoi_phai(O, hinh, ws, h, gap, canh_bao):
    """REDPUMP: ô tên model ở CUỐI PHẢI mỗi đường. Ô cùng cột dàn dọc theo thứ tự cao độ điểm cuối
    (lệch > nửa ô thì có đường dẫn ngắn). Chỉ ô bị vướng mới dời, và vẫn sát cuối đường."""
    n = len(hinh)
    cuoi = [tuple(g['A'][-1]) for g in hinh]
    kq = [None] * n
    nhom, cur, hi = [], [], -1e9
    for k in sorted(range(n), key=lambda k: cuoi[k][0]):
        lo = cuoi[k][0] + gap
        if cur and lo > hi:
            nhom.append(cur)
            cur, hi = [], -1e9
        cur.append(k)
        hi = max(hi, lo + ws[k] + 6)
    if cur:
        nhom.append(cur)
    hong = []
    for nh in nhom:
        nh = sorted(nh, key=lambda k: cuoi[k][1])
        tam = _dan_doc([cuoi[k][1] for k in nh], h, 2.0, O.P.y0 + 1.5, O.P.y1 - 1.5)
        cot = any(abs(c - cuoi[k][1]) > 0.5 * h for k, c in zip(nh, tam))   # có ô phải dời -> cả cột thẳng hàng + đường dẫn
        for k, c in zip(nh, tam):
            r = _thu_o_phai(O, k, cuoi[k], ws[k], h, gap, c, cot)
            if r:
                O.nhan.append(r[0])
                kq[k] = r
            else:
                hong.append(k)
    for k in hong:
        r = None
        for j in range(1, 8):
            for dau in (1, -1):
                r = _thu_o_phai(O, k, cuoi[k], ws[k], h, gap, cuoi[k][1] + dau * j * 0.55 * h)
                if r:
                    break
            if r:
                break
        if r is None:
            r0 = dat_phai_gan(O, hinh[k]['A'], ws[k], h, gap)
            r = (r0, None) if r0 else None
        if r is None:
            r0, dan = dat_dan(O, k, hinh[k]['A'], ws[k], h, (1.0, 0.97, 0.94, 0.9, 0.85))
            if r0 is not None:
                r = (r0, dan)
                canh_bao.append(f'Nhãn {hinh[k]["md"]["khoa"]}: cuối đường không đủ chỗ — đặt gần cuối, có đường dẫn.')
        if r is None:
            raise LoiBoCuc(f'Không đặt được nhãn {hinh[k]["md"]["khoa"]} ở cuối đường mà không đè — cần tách biểu đồ '
                           f'hoặc tăng --cao-pt.')
        O.nhan.append(r[0])
        kq[k] = r
    return kq


def dat_tron(O, i, A, D):
    """AQUARIS: số khoanh tròn nằm trên chính đường của model."""
    xs = A[:, 0]
    x_s, x_e = xs[0], xs[-1]
    for t in (0.72, 0.62, 0.82, 0.52, 0.9, 0.42, 0.32, 0.95, 0.22, 0.12,
              0.67, 0.57, 0.77, 0.47, 0.87, 0.37, 0.27, 0.17, 0.07):
        k = int(np.argmin(np.abs(xs - (x_s + t * (x_e - x_s)))))
        cx, cy = A[k]
        r = pymupdf.Rect(cx - D / 2, cy - D / 2, cx + D / 2, cy + D / 2)
        if O.tu_do(r, bo_qua=(i,), dem_n=2.0):
            return r
    return None


def _cat_doan(Pts, a, b, dem):
    """Có điểm nào của Pts cách đoạn a–b dưới dem pt không."""
    a, b = np.asarray(a, float), np.asarray(b, float)
    ab = b - a
    L2 = float(ab @ ab)
    if L2 < 1e-9:
        return bool(np.any(np.hypot(*(Pts - a).T) < dem))
    t = np.clip(((Pts - a) @ ab) / L2, 0.0, 1.0)
    gan = a + t[:, None] * ab
    return bool(np.any(np.hypot(*(Pts - gan).T) < dem))


def dat_dan(O, i, A, w, h, phan_so):
    """Dự phòng khi sát đường không đủ chỗ: nhãn ở chỗ trống gần nhất + đường dẫn mảnh từ một điểm trên đường.
    Ưu tiên đường dẫn ngắn, KHÔNG cắt đường khác, không qua nhãn khác; neo theo thứ tự phan_so của kiểu hãng."""
    n = len(A)
    khac = [B for k, B in O.duong if k != i]
    Pk = np.vstack(khac) if khac else np.zeros((0, 2))
    Ek = np.array([B[j] for B in khac for j in (0, -1)]) if khac else np.zeros((0, 2))   # đầu/cuối đường khác
    du_phong = None
    for rad in (10, 14, 18, 24, 30, 38, 48, 60, 75, 95, 120):
        for f in phan_so:
            ax, ay = A[min(n - 1, int(round(f * (n - 1))))]
            if len(Pk) and np.min(np.hypot(Pk[:, 0] - ax, Pk[:, 1] - ay)) < 3.0:
                continue                       # neo sát đường khác -> dễ hiểu nhầm
            for goc in range(0, 360, 15):
                cx = ax + (rad + w / 2) * math.cos(math.radians(goc))
                cy = ay + (rad + h / 2) * math.sin(math.radians(goc))
                r = pymupdf.Rect(cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2)
                if not O.tu_do(r, dem_n=2.0):
                    continue
                px, py = min(max(ax, r.x0), r.x1), min(max(ay, r.y0), r.y1)
                L = math.hypot(px - ax, py - ay)
                if L < 3:
                    continue
                ex, ey = px - (px - ax) / L * 1.5, py - (py - ay) / L * 1.5
                qua_nhan = any(_cat_doan(np.array([[q.x0, q.y0], [q.x1, q.y0], [q.x0, q.y1], [q.x1, q.y1],
                                                   [(q.x0 + q.x1) / 2, (q.y0 + q.y1) / 2]]), (ax, ay), (ex, ey), 1.0)
                               or (q.x0 <= (ax + ex) / 2 <= q.x1 and q.y0 <= (ay + ey) / 2 <= q.y1) for q in O.nhan)
                if qua_nhan:
                    continue
                if len(Pk) == 0 or not (_cat_doan(Pk, (ax, ay), (ex, ey), 2.5) or _cat_doan(Ek, (ax, ay), (ex, ey), 6.0)):
                    return r, (ax, ay, ex, ey)
                if du_phong is None:
                    du_phong = (r, (ax, ay, ex, ey))
    return du_phong if du_phong else (None, None)


# ----------------------------------------------------------------------------- biểu đồ
def do_bieu_do_co_dinh(S, ph, W, ds):
    """Phần cao cố định (không kể vùng vẽ) của khối biểu đồ: trên, dưới, chú giải, khối điều kiện."""
    co = S.co_so
    hc = ph.cao(co)
    off = hc / 2.0 + 2.0
    tren = (2 * hc + 8) if S.bt['lmin'] == 'tren' else 4.0
    if S.bt['lmin'] == 'duoi':
        duoi = off + hc + 5 + 4 + hc + 3
    else:
        duoi = off + hc + 3 + hc + 2
    chu_giai = 0.0
    if S.kieu_nhan == 'so_khoanh_tron_va_chu_giai':
        D = co + 5
        w_muc = D + 4 + max(ph.rong(md['khoa'], co) for md in ds) + 14
        ncot = max(1, int((W - 8) // w_muc))
        chu_giai = 6 + math.ceil(len(ds) / ncot) * (D + 3)
    return tren, duoi, chu_giai


def dong_chu(ph, s, co, w):
    tu, dong, cur = s.split(' '), [], ''
    for t in tu:
        thu = (cur + ' ' + t).strip()
        if ph.rong(thu, co) <= w or not cur:
            cur = thu
        else:
            dong.append(cur)
            cur = t
    if cur:
        dong.append(cur)
    return dong


def ve_bieu_do(K, S, R, ds, dk, canh_bao):
    ph, co = K.ph, S.co_so
    hc = ph.cao(co)
    off = hc / 2.0 + 2.0
    bt = S.bt
    tren, duoi, chu_giai = do_bieu_do_co_dinh(S, ph, R.width, ds)
    dong_dk = dong_chu(ph, dk, S.co_than, R.width)
    cao_dk = 6 + len(dong_dk) * S.co_than * 1.25
    ph_h = R.height - tren - duoi - chu_giai - cao_dk
    if ph_h < 110:
        raise LoiBoCuc(f'Khung biểu đồ còn {ph_h:.0f} pt chiều cao vùng vẽ (< 110 pt) — cần tăng --cao-pt hoặc tách bảng.')
    y0 = R.y0 + tren
    y1 = y0 + ph_h

    def rong_so(vals):
        return max(ph.rong(fmt(v), co) for v in vals)

    # ---- trục H
    h_max = max(p[1] for md in ds for p in md['diem'])
    h_eff = h_max / max(0.5, 1.0 - (hc + 7) / ph_h)
    bh, top_h = chon_truc(h_eff, ph_h, lambda v: hc, max(2.2 * hc, 22.0))
    ticks_h = [i * bh for i in range(int(round(top_h / bh)) + 1)]
    w_nhan_h = rong_so(ticks_h)
    le_h = 3 + 3 + w_nhan_h + 5 + hc + 1
    le_khac = max(6.0, ph.rong(fmt(0), co) / 2 + 2)
    if bt['lmin'] == 'duoi':
        w_dv = max(ph.rong('Q (m³/h)', co), ph.rong('Q (l/min)', co)) + 8
        le_h = max(le_h, w_dv)
    if S.truc_h == 'trai':
        x0, x1 = R.x0 + le_h, R.x1 - 10
    else:
        x0, x1 = R.x0 + le_khac + 8, R.x1 - le_h
    pw = x1 - x0
    # ---- trục Q (REDPUMP chừa chỗ cho ô tên ở cuối đường)
    q_max = max(p[0] for md in ds for p in md['diem'])
    q_eff = q_max
    pad_x, pad_y = 3.0, 1.6
    if S.kieu_nhan == 'ten_model_trong_o_mau_cuoi_phai_duong':
        can = max(ph.rong(md['khoa'], co, True) for md in ds) + 2 * pad_x + S.net + 4 + 6
        if can / pw < 0.35:
            q_eff = q_max / (1 - can / pw)
    bq, top_q = chon_truc(q_eff, pw, rong_so, 26.0)
    ticks_q = [i * bq for i in range(int(round(top_q / bq)) + 1)]
    top_l = top_q * LMIN
    bl = chon_buoc_phu(top_l, pw / top_l, rong_so, 26.0)
    ticks_l = [i * bl for i in range(int(math.floor(top_l / bl + 1e-9)) + 1)]
    P = pymupdf.Rect(x0, y0, x1, y1)

    def X(q):
        return x0 + q / top_q * pw

    def Y(h):
        return y1 - h / top_h * ph_h

    # ---- đường (hình học trước, vẽ sau)
    mau_ds = mau_duong(S, len(ds), canh_bao)
    O = ODat(P, S.net / 2 + 0.9)
    hinh = []
    for i, md in enumerate(ds):
        q0, qn = md['diem'][0][0], md['diem'][-1][0]
        qd = np.linspace(q0, qn, 3000)
        A = np.column_stack([X(qd), Y(md['f'](qd))])
        qv = np.union1d(np.linspace(q0, qn, 220), [p[0] for p in md['diem']])
        B = np.column_stack([X(qv), Y(md['f'](qv))])
        O.duong.append((i, A))
        hinh.append(dict(md=md, A=A, B=B, mau=mau_ds[i]))
    # ---- đặt nhãn
    nhan = [None] * len(ds)
    gap = S.net / 2 + 1.4
    if S.kieu_nhan == 'so_khoanh_tron_va_chu_giai':
        D = co + 5
        for i, g in enumerate(hinh):
            r = dat_tron(O, i, g['A'], D)
            dan = None
            if r is None:
                r, dan = dat_dan(O, i, g['A'], D, D, (0.7, 0.5, 0.85, 0.3, 1.0, 0.15, 0.0))
                if r is not None:
                    canh_bao.append(f'Số {i + 1} ({g["md"]["khoa"]}): không đủ chỗ trên đường — đặt gần, có đường dẫn.')
            if r is None:
                raise LoiBoCuc(f'Không đặt được số khoanh tròn cho {g["md"]["khoa"]} mà không đè — cần tách biểu đồ.')
            O.nhan.append(r)
            nhan[i] = (r, dan)
    elif S.kieu_nhan == 'ten_model_trong_o_mau_cuoi_phai_duong':
        ws = [ph.rong(g['md']['khoa'], co, True) + 2 * pad_x for g in hinh]
        nhan = dat_cuoi_phai(O, hinh, ws, hc + 2 * pad_y, gap, canh_bao)
    else:
        thu_tu = sorted(range(len(ds)), key=lambda k: hinh[k]['A'][0, 1])
        for i in thu_tu:
            g = hinh[i]
            w, h = ph.rong(g['md']['khoa'], co), hc
            r = dat_trai(O, g['A'], w, h, gap)
            dan = None
            if r is None:
                r, dan = dat_dan(O, i, g['A'], w, h, (0.0, 0.1, 0.2, 0.3, 0.45, 0.6, 0.75, 0.9, 1.0))
                if r is not None:
                    canh_bao.append(f'Nhãn {g["md"]["khoa"]}: không đủ chỗ sát đường — đặt gần, có đường dẫn.')
            if r is None:
                raise LoiBoCuc(f'Không đặt được nhãn {g["md"]["khoa"]} mà không đè — cần tách biểu đồ.')
            O.nhan.append(r)
            nhan[i] = (r, dan)
    # ---- nền, lưới, khung
    pg = K.page
    if do_choi(S.nen_bd) < 0.999:
        pg.draw_rect(P, color=None, fill=S.nen_bd, width=0)
    for q in ticks_q:
        if 0 < q < top_q or bt['khung'] == 'duoi':
            pg.draw_line((X(q), y0), (X(q), y1), color=S.luoi_mau, width=S.luoi)
    for h in ticks_h:
        if 0 < h < top_h or bt['khung'] == 'duoi':
            pg.draw_line((x0, Y(h)), (x1, Y(h)), color=S.luoi_mau, width=S.luoi)
    tr = S.chu
    if bt['khung'] == 'day_du':
        pg.draw_rect(P, color=tr, width=0.75)
    elif bt['khung'] == 'trai_duoi':
        pg.draw_line((x0, y0), (x0, y1), color=tr, width=0.75)
        pg.draw_line((x0, y1), (x1, y1), color=tr, width=0.75)
    else:
        pg.draw_line((x0, y1), (x1, y1), color=tr, width=0.75)
    # vạch + số trục Q (dưới)
    nen_ngoai = TRANG
    for q in ticks_q:
        pg.draw_line((X(q), y1), (X(q), y1 + 3), color=tr, width=0.5)
        K.chu(X(q), y1 + off + ph.asc[False] * co, fmt(q), co, tr, nen_ngoai, canh='giua', loai='truc')
    y_hang1 = y1 + off + hc
    # trục H
    xh = x0 if S.truc_h == 'trai' else x1
    huong = -1 if S.truc_h == 'trai' else 1
    if bt['khung'] == 'duoi':
        pg.draw_line((xh, y0), (xh, y1), color=tr, width=0.75)
    for h in ticks_h:
        pg.draw_line((xh, Y(h)), (xh + 3 * huong, Y(h)), color=tr, width=0.5)
        K.chu(xh + 6 * huong, ph.chan(Y(h), co), fmt(h), co, tr, nen_ngoai,
              canh='phai' if huong < 0 else 'trai', loai='truc')
    x_tieu_h = xh + huong * (6 + w_nhan_h + 5 + hc / 2)
    K.chu(x_tieu_h, (y0 + y1) / 2, 'H (m)', co, tr, nen_ngoai, dam=True, xoay=90, loai='truc')
    # thang l/min phụ, song song
    if bt['lmin'] == 'duoi':
        K.chu(x0 - 6, y1 + off + ph.asc[False] * co, 'Q (m³/h)', co, tr, nen_ngoai, dam=True, canh='phai', loai='truc')
        ya = y_hang1 + 5
        pg.draw_line((x0, ya), (x1, ya), color=tr, width=0.5)
        for v in ticks_l:
            xv = x0 + v / top_l * pw
            pg.draw_line((xv, ya), (xv, ya + 3), color=tr, width=0.5)
            K.chu(xv, ya + 4 + ph.asc[False] * co, fmt(v), co, tr, nen_ngoai, canh='giua', loai='truc')
        K.chu(x0 - 6, ya + 4 + ph.asc[False] * co, 'Q (l/min)', co, tr, nen_ngoai, dam=True, canh='phai', loai='truc')
        y_cuoi = ya + 4 + hc + 3
    else:
        K.chu((x0 + x1) / 2, y_hang1 + 3 + ph.asc[True] * co, 'Q (m³/h)', co, tr, nen_ngoai, dam=True,
              canh='giua', loai='truc')
        y_cuoi = y_hang1 + 3 + hc + 2
        if bt['khung'] != 'day_du':
            pg.draw_line((x0, y0), (x1, y0), color=tron(tr, TRANG, 0.5) if bt['khung'] == 'duoi' else tr,
                         width=0.5)
        for v in ticks_l:
            xv = x0 + v / top_l * pw
            pg.draw_line((xv, y0), (xv, y0 - 3), color=tr, width=0.5)
            K.chu(xv, y0 - off + ph.desc[False] * co, fmt(v), co, tr, nen_ngoai, canh='giua', loai='truc')
        K.chu((x0 + x1) / 2, y0 - off - hc - 2 + ph.desc[True] * co, 'Q (l/min)', co, tr, nen_ngoai,
              dam=True, canh='giua', loai='truc')
    # ---- đường cong (có chừa khe dưới số khoanh tròn)
    for i, g in enumerate(hinh):
        B = g['B']
        if S.kieu_nhan == 'so_khoanh_tron_va_chu_giai':
            r = nhan[i][0]
            c = ((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
            rad = r.width / 2 + 0.6
            Bd = g['A'][::6]
            Bd = np.vstack([Bd, g['A'][-1]])
            ngoai = np.hypot(Bd[:, 0] - c[0], Bd[:, 1] - c[1]) > rad
            doan, cur = [], []
            for p, ok in zip(Bd, ngoai):
                if ok:
                    cur.append(p)
                elif cur:
                    doan.append(cur)
                    cur = []
            if cur:
                doan.append(cur)
        else:
            doan = [list(B)]
        for dd in doan:
            if len(dd) >= 2:
                pg.draw_polyline([tuple(p) for p in dd], color=g['mau'], width=S.net, lineCap=1, lineJoin=1)
    # ---- nhãn
    ds_nhan = []
    for i, g in enumerate(hinh):
        r, dan = nhan[i]
        ten = g['md']['khoa']
        if dan:
            pg.draw_line((dan[0], dan[1]), (dan[2], dan[3]), color=g['mau'], width=0.6)
            if (dan[0], dan[1]) != tuple(g['A'][-1]):          # neo giữa đường: chấm đánh dấu
                pg.draw_circle((dan[0], dan[1]), 1.6, color=None, fill=g['mau'], width=0)
        if S.kieu_nhan == 'so_khoanh_tron_va_chu_giai':
            c = ((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
            ve_so_tron(K, c, r.width / 2, i + 1, g['mau'], co)
        elif S.kieu_nhan == 'ten_model_dau_trai_duong':
            K.chu(r.x0, r.y0 + ph.asc[False] * co, ten, co, g['mau'], S.nen_bd, loai='nhan_duong')
        else:
            pg.draw_rect(r, color=None, fill=g['mau'], width=0)
            K.chu(r.x0 + pad_x, ph.chan((r.y0 + r.y1) / 2, co, True), ten, co, TRANG, g['mau'], dam=True,
                  loai='nhan_duong')
        ds_nhan.append((ten, tuple(round(v, 2) for v in r)))
    # ---- chú giải số -> model
    y = y_cuoi
    if S.kieu_nhan == 'so_khoanh_tron_va_chu_giai':
        D = co + 5
        w_muc = D + 4 + max(ph.rong(md['khoa'], co) for md in ds) + 14
        ncot = max(1, int((R.width - 8) // w_muc))
        y += 6
        for i, g in enumerate(hinh):
            hang, cot = divmod(i, ncot)
            cx = R.x0 + 4 + cot * w_muc + D / 2
            cy = y + hang * (D + 3) + D / 2
            ve_so_tron(K, (cx, cy), D / 2, i + 1, g['mau'], co)
            K.chu(cx + D / 2 + 4, ph.chan(cy, co), g['md']['khoa'], co, tr, TRANG, loai='chu_giai')
        y += math.ceil(len(ds) / ncot) * (D + 3)
    # ---- khối điều kiện
    y += 6
    for d in dong_dk:
        K.chu(R.x0, y + ph.asc[False] * S.co_than, d, S.co_than, tr, TRANG, loai='dieu_kien')
        y += S.co_than * 1.25
    return dict(nhan=ds_nhan, mau={g['md']['khoa']: hexm(g['mau']) for g in hinh},
                cuoi={g['md']['khoa']: tuple(round(float(v), 2) for v in g['A'][-1]) for g in hinh},
                truc=dict(Q=(0, top_q, bq), H=(0, top_h, bh), lmin=(0, ticks_l[-1], bl)), vung_ve=tuple(P))


def ve_so_tron(K, c, rad, so, mau_nen, co):
    K.page.draw_circle(c, rad, color=TRANG, fill=mau_nen, width=0.8)
    mau_so = max((TRANG, DEN), key=lambda m: tuong_phan(m, mau_nen))
    K.chu(c[0], K.ph.chan(c[1], co, True), str(so), co, mau_so, mau_nen, dam=True, canh='giua', loai='so_tron')


# ----------------------------------------------------------------------------- bảng
def _h_tai(md, q):
    for p in md['diem']:
        if abs(p[0] - q) < 1e-9:
            return p[1]
    return None


def _o_h(md, q):
    """Chữ ô H: Q ngoài dải chạy của model -> "–" (QTH 4.7, CTB 3.2); trong dải mà thiếu số -> trống (lộ lỗi dữ liệu)."""
    qs = [p[0] for p in md['diem']]
    if qs and (q < min(qs) - 1e-9 or q > max(qs) + 1e-9):
        return '–'
    return fmt(_h_tai(md, q))


def dung_bang(S, ph, ds, W, canh_bao, bo_cot_a=False):
    """Dựng lưới bảng (chưa vẽ). Thử cỡ bang_nen rồi bang_min; không vừa thì báo cần tách."""
    cot_q = sorted({p[0] for md in ds for p in md['diem']})
    for md in ds:
        if md['diem']:
            qa, qb = md['diem'][0][0], md['diem'][-1][0]
            thieu = [q for q in cot_q if qa < q < qb and _h_tai(md, q) is None]
            if thieu:
                canh_bao.append(f'{md["khoa"]}: ô trống giữa dãy H tại Q = {", ".join(fmt(q) for q in thieu)} m³/h '
                                f'(nguồn không có điểm, không tự nội suy).')
    co_1p = any(md['m1'] for md in ds)
    co_3p = any(md['m3'] for md in ds)
    co_a = any(md['a1'] is not None or md['a3'] is not None for md in ds)
    hai_a = any(md['a1'] is not None and md['a3'] is not None for md in ds)
    if not co_1p:
        canh_bao.append('Không model nào có bản 1 pha — bỏ cột "Model 1~".')
    if not co_a:
        canh_bao.append('Nguồn không có dòng điện A — bỏ cột A (QĐ 152).')
    else:
        thieu = [md['khoa'] for md in ds if md['a1'] is None and md['a3'] is None]
        if thieu:
            canh_bao.append(f'THIẾU DÒNG ĐIỆN A ở {len(thieu)} model ({", ".join(thieu)}) — giữ cột A, ô để trống; '
                            f'điền số suy ra vào DU_LIEU theo QTH 4.6 (QĐ 73) rồi vẽ lại. (Không tự suy, không điền 0.)')
    for md in ds:
        if md['kw'] is None or md['hp'] is None:
            canh_bao.append(f'{md["khoa"]}: thiếu P2 kW/HP — để trống ô.')

    def a_txt(md):
        if md['a1'] is not None and md['a3'] is not None:
            return f'{fmt(md["a1"])} / {fmt(md["a3"])}'
        return fmt(md['a3'] if md['a3'] is not None else md['a1'])

    # mô tả cột thuộc tính: (nhãn trên, nhãn dưới, hàm lấy giá trị, canh, đậm)
    tt = []
    if co_1p:
        tt.append(('Model', '1~', lambda md: md['m1'] or '–', 'trai', False))
    if co_3p:
        tt.append(('Model', '3~', lambda md: md['m3'], 'trai', True))
    tt.append(('P2', 'kW', lambda md: fmt(md['kw']), 'phai', False))
    tt.append(('P2', 'HP', lambda md: fmt(md['hp']), 'phai', False))
    if co_a:
        tt.append(('I', 'A 1~ / 3~' if hai_a else 'A', a_txt, 'phai', False))
    dv = ('Q (m³/h)', 'Q (l/min)', 'H (m)')
    for co in (S.co_so, S.co_min_bang):
        pad = 4.0
        hh = ph.cao(co) + 2 * 3.2
        if S.bang_huong == 'model_theo_dong':
            w_tt = [max([ph.rong(a, co, True), ph.rong(b, co, True)] +
                        [ph.rong(f(md), co, dam) for md in ds]) + 2 * pad for a, b, f, _, dam in tt]
            w_dv = max(ph.rong(s, co, True) for s in dv) + 2 * pad
            w_q = max([ph.rong(fmt(q), co, True) for q in cot_q] + [ph.rong(fmt(round(q * LMIN)), co) for q in cot_q] +
                      [ph.rong(_o_h(md, q), co) for md in ds for q in cot_q]) + 2 * pad
            tong = sum(w_tt) + w_dv + w_q * len(cot_q)
            if tong <= W:
                k = min(GIAN_BANG_TOI_DA, W / tong)      # giãn đều mọi cột, tối đa 3 lần
                ws = [w * k for w in list(w_tt) + [w_dv] + [w_q] * len(cot_q)]
                return dict(huong='dong', co=co, cot_w=ws, cao_hang=hh, so_hang=2 + len(ds), tt=tt, cot_q=cot_q,
                            ds=ds, cao=hh * (2 + len(ds)), rong=sum(ws), dv=dv)
        else:
            nh = [(a, b, f, c, dam) for a, b, f, c, dam in tt]
            w0 = max([ph.rong(a, co, True) for a, *_ in nh] + [ph.rong(dv[0], co, True)] +
                     [ph.rong(fmt(q), co) for q in cot_q]) + 2 * pad
            w1 = max([ph.rong(b, co, True) for _, b, *_ in nh] + [ph.rong(dv[1], co, True)] +
                     [ph.rong(fmt(round(q * LMIN)), co) for q in cot_q]) + 2 * pad
            w_m = max([ph.rong(f(md), co, dam) for _, _, f, _, dam in nh for md in ds] +
                      [ph.rong(_o_h(md, q), co) for md in ds for q in cot_q] + [co + 5]) + 2 * pad
            tong = w0 + w1 + w_m * len(ds)
            if tong <= W:
                k = min(GIAN_BANG_TOI_DA, W / tong)
                ws = [w * k for w in [w0, w1] + [w_m] * len(ds)]
                so_hang = 1 + len(nh) + 1 + len(cot_q)
                hang_tron = co + 5 + 4
                return dict(huong='cot', co=co, cot_w=ws, cao_hang=hh, so_hang=so_hang, tt=nh, cot_q=cot_q, ds=ds,
                            cao=hang_tron + hh * (so_hang - 1), rong=sum(ws), dv=dv, hang_tron=hang_tron)
    raise LoiBoCuc(f'Bảng {len(ds)} model × {len(cot_q)} cột Q rộng hơn {W:.0f} pt ngay cả ở cỡ {S.co_min_bang:g} pt '
                   f'— cần tách bảng (không thu chữ dưới ngưỡng).')


def ve_bang(K, S, B, x0, y0, mau_so=None):
    ph, pg, co = K.ph, K.page, B['co']
    ws, hh = B['cot_w'], B['cao_hang']
    xs = [x0]
    for w in ws:
        xs.append(xs[-1] + w)
    chu = S.chu
    ke = S.bang_ke
    nhat = tron(chu, TRANG, 0.62)
    if ke == 'dau_bang_toi_chu_trang_soc_xam':
        nen_dau, chu_dau, soc = S.chu, TRANG, S.m.get('nen_bieu_do', S.nen_bd)
    elif ke == 'day_du_manh':
        nen_dau, chu_dau, soc = S.m.get('nen', TRANG), chu, None
    else:
        nen_dau, chu_dau, soc = TRANG, chu, None
    nen_nhan = S.m.get('nen', TRANG) if ke == 'ngang_manh' else nen_dau

    def o(c0, c1, r0, r1, ys):
        return pymupdf.Rect(xs[c0], ys[r0], xs[c1], ys[r1])

    if B['huong'] == 'dong':
        n_tt = len(B['tt'])
        ys = [y0 + k * hh for k in range(B['so_hang'] + 1)]
        nc = len(ws)
        R_all = pymupdf.Rect(xs[0], ys[0], xs[-1], ys[-1])
        # nền
        pg.draw_rect(o(0, nc, 0, 2, ys), color=None, fill=nen_dau, width=0)
        for k, md in enumerate(B['ds']):
            if soc is not None and k % 2 == 1:
                pg.draw_rect(o(0, nc, 2 + k, 3 + k, ys), color=None, fill=soc, width=0)
        # đầu bảng
        for c, (a, b, f, canh, dam) in enumerate(B['tt']):
            if a == 'Model':
                K.chu_giua_o(o(c, c + 1, 0, 2, ys), f'{a} {b}', co, chu_dau, nen_dau, True, 'trai')
            else:
                K.chu_giua_o(o(c, c + 1, 0, 1, ys), a, co, chu_dau, nen_dau, True, 'phai')
                K.chu_giua_o(o(c, c + 1, 1, 2, ys), b, co, chu_dau, nen_dau, False, 'phai')
        cdv = n_tt
        K.chu_giua_o(o(cdv, cdv + 1, 0, 1, ys), B['dv'][0], co, chu_dau, nen_dau, True, 'trai')
        K.chu_giua_o(o(cdv, cdv + 1, 1, 2, ys), B['dv'][1], co, chu_dau, nen_dau, True, 'trai')
        for j, q in enumerate(B['cot_q']):
            c = n_tt + 1 + j
            K.chu_giua_o(o(c, c + 1, 0, 1, ys), fmt(q), co, chu_dau, nen_dau, True, 'phai')
            K.chu_giua_o(o(c, c + 1, 1, 2, ys), fmt(round(q * LMIN)), co, chu_dau, nen_dau, False, 'phai')
        # thân
        if soc is not None:      # cột đơn vị gộp: nền trắng
            pg.draw_rect(o(cdv, cdv + 1, 2, B['so_hang'], ys), color=None, fill=TRANG, width=0)
        for k, md in enumerate(B['ds']):
            nen = soc if (soc is not None and k % 2 == 1) else TRANG
            for c, (a, b, f, canh, dam) in enumerate(B['tt']):
                K.chu_giua_o(o(c, c + 1, 2 + k, 3 + k, ys), f(md), co, chu, nen, dam, canh)
            for j, q in enumerate(B['cot_q']):
                c = n_tt + 1 + j
                K.chu_giua_o(o(c, c + 1, 2 + k, 3 + k, ys), _o_h(md, q), co, chu, nen, False, 'phai')
        K.chu_giua_o(o(cdv, cdv + 1, 2, B['so_hang'], ys), B['dv'][2], co, chu, TRANG, True, 'trai')
        # kẻ
        if ke == 'day_du_manh':
            for k in range(B['so_hang'] + 1):
                if k == 1:
                    for c in range(n_tt, nc):
                        pg.draw_line((xs[c], ys[1]), (xs[c + 1], ys[1]), color=nhat, width=0.5)
                    continue
                if k > 2:
                    pg.draw_line((xs[0], ys[k]), (xs[cdv], ys[k]), color=nhat, width=0.5)
                    pg.draw_line((xs[cdv + 1], ys[k]), (xs[-1], ys[k]), color=nhat, width=0.5)
                else:
                    pg.draw_line((xs[0], ys[k]), (xs[-1], ys[k]), color=nhat, width=0.5)
            for c in range(1, nc):
                pg.draw_line((xs[c], ys[0]), (xs[c], ys[-1]), color=nhat, width=0.5)
            pg.draw_rect(R_all, color=chu, width=0.8)
            pg.draw_line((xs[0], ys[2]), (xs[-1], ys[2]), color=chu, width=0.8)
        elif ke == 'dau_bang_toi_chu_trang_soc_xam':
            pg.draw_line((xs[0], ys[2]), (xs[-1], ys[2]), color=S.m.get('nhan', chu), width=1.2)
            pg.draw_line((xs[0], ys[-1]), (xs[-1], ys[-1]), color=chu, width=0.6)
            pg.draw_line((xs[cdv], ys[2]), (xs[cdv], ys[-1]), color=nhat, width=0.4)
            pg.draw_line((xs[cdv + 1], ys[2]), (xs[cdv + 1], ys[-1]), color=nhat, width=0.4)
        else:
            for k in range(3, B['so_hang']):
                pg.draw_line((xs[0], ys[k]), (xs[-1], ys[k]), color=nhat, width=0.4)
            pg.draw_line((xs[0], ys[0]), (xs[-1], ys[0]), color=chu, width=0.9)
            pg.draw_line((xs[0], ys[2]), (xs[-1], ys[2]), color=chu, width=0.6)
            pg.draw_line((xs[0], ys[-1]), (xs[-1], ys[-1]), color=chu, width=0.9)
        return
    # ---- mỗi model một cột (AQUARIS)
    nm = len(B['ds'])
    ht = B['hang_tron']
    ys = [y0, y0 + ht] + [y0 + ht + k * hh for k in range(1, B['so_hang'])]
    nc = len(ws)
    n_tt = len(B['tt'])
    r_dv = 1 + n_tt                      # hàng tiêu đề phụ Q / H
    pg.draw_rect(o(0, 2, 1, B['so_hang'], ys), color=None, fill=nen_nhan, width=0)
    for j, md in enumerate(B['ds']):
        r = o(2 + j, 3 + j, 0, 1, ys)
        c = ((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
        ve_so_tron(K, c, (co + 5) / 2, j + 1, mau_so[j] if mau_so else S.chu, co)
    for t, (a, b, f, canh, dam) in enumerate(B['tt']):
        rr = 1 + t
        K.chu_giua_o(o(0, 1, rr, rr + 1, ys), a, co, chu, nen_nhan, True, 'trai')
        K.chu_giua_o(o(1, 2, rr, rr + 1, ys), b, co, chu, nen_nhan, False, 'trai')
        for j, md in enumerate(B['ds']):
            K.chu_giua_o(o(2 + j, 3 + j, rr, rr + 1, ys), f(md), co, chu, TRANG, dam, 'phai')
    K.chu_giua_o(o(0, 1, r_dv, r_dv + 1, ys), B['dv'][0], co, chu, nen_nhan, True, 'phai')
    K.chu_giua_o(o(1, 2, r_dv, r_dv + 1, ys), B['dv'][1], co, chu, nen_nhan, True, 'phai')
    K.chu_giua_o(o(2, nc, r_dv, r_dv + 1, ys), B['dv'][2], co, chu, TRANG, True, 'giua')
    for k, q in enumerate(B['cot_q']):
        rr = r_dv + 1 + k
        K.chu_giua_o(o(0, 1, rr, rr + 1, ys), fmt(q), co, chu, nen_nhan, False, 'phai')
        K.chu_giua_o(o(1, 2, rr, rr + 1, ys), fmt(round(q * LMIN)), co, chu, nen_nhan, False, 'phai')
        for j, md in enumerate(B['ds']):
            K.chu_giua_o(o(2 + j, 3 + j, rr, rr + 1, ys), _o_h(md, q), co, chu, TRANG, False, 'phai')
    for k in range(2, B['so_hang']):
        if k in (r_dv, r_dv + 1):
            continue
        pg.draw_line((xs[0], ys[k]), (xs[-1], ys[k]), color=nhat, width=0.4)
    pg.draw_line((xs[0], ys[1]), (xs[-1], ys[1]), color=chu, width=0.9)
    pg.draw_line((xs[0], ys[r_dv]), (xs[-1], ys[r_dv]), color=chu, width=0.6)
    pg.draw_line((xs[0], ys[r_dv + 1]), (xs[-1], ys[r_dv + 1]), color=chu, width=0.6)
    pg.draw_line((xs[0], ys[-1]), (xs[-1], ys[-1]), color=chu, width=0.9)


# ----------------------------------------------------------------------------- dựng trang
def chuoi_dien(S, ds):
    s = []
    if any(md['m1'] for md in ds):
        s.append(S.dien['mot_pha'])
    if any(md['m3'] for md in ds):
        s.append(S.dien['ba_pha'])
    return '  ·  '.join(s)


def ghi_canh_bao(ra, dong, tieu_de):
    """Ghi <ra>.canh_bao.txt (ghi đè) để người dựng / điều phối đọc lại."""
    tep = Path(str(ra) + '.canh_bao.txt')
    try:
        tep.parent.mkdir(parents=True, exist_ok=True)
        noi_dung = tieu_de + '\n' + ('\n'.join(f'- {c}' for c in dong) if dong else '- Không có cảnh báo.') + '\n'
        tep.write_text(noi_dung, encoding='utf-8')
    except OSError:
        return None
    return str(tep)


def tep_goc(du_lieu, so_voi=None):
    """Tệp số liệu gốc đã duyệt để so khi gặp gợn: --so-voi; không có thì DU_LIEU/<cùng tên> (tệp vào ngoài
    DU_LIEU/) hoặc chính tệp vào (tệp vào trong DU_LIEU/). Không có bản gốc -> None."""
    if so_voi:
        p = Path(so_voi)
        if not p.is_file():
            raise LoiVe(f'Không thấy tệp --so-voi: {p}')
        return p
    vao = Path(du_lieu).resolve()
    goc_dl = Path(THU_MUC_DU_LIEU).resolve()
    if goc_dl in vao.parents:
        return vao
    p = goc_dl / vao.name
    return p if p.is_file() else None


def ve(du_lieu, hang, ra, models=None, kieu='ca_hai', rong=None, cao=None, ho_so=None, tieu_de=True,
       cho_phep_gan_giong=False, chap_nhan_gon=False, bo_cot_a=False, so_voi=None):
    """Vẽ và lưu PDF + <ra>.canh_bao.txt. Trả về báo cáo (dict). Lỗi: LoiDuLieu (2) / LoiBoCuc (3) / LoiVe (1)."""
    if kieu not in ('bieu_do', 'bang', 'ca_hai'):
        raise LoiVe(f'--kieu "{kieu}" không hợp lệ (bieu_do | bang | ca_hai).')
    hs_path = Path(ho_so) if ho_so else HO_SO_MAC_DINH
    if not hs_path.is_file():
        raise LoiVe(f'Không thấy hồ sơ hãng: {hs_path} (dùng --ho-so).')
    HS = json.loads(hs_path.read_text(encoding='utf-8'))
    hang = str(hang).strip().upper()
    if hang not in HS['hang']:
        raise LoiVe(f'Hãng "{hang}" không có trong hồ sơ ({", ".join(HS["hang"])}).')
    S = KieuHang(HS, hang)
    tt, ds_all, theo_ten = doc_du_lieu(du_lieu, hang)
    ds = chon_model(ds_all, theo_ten, models, hang)
    canh_bao = []
    gon = {}
    for md in ds:
        md['f'] = kiem_duong(md, can_duong=(kieu != 'bang'))
        if kieu != 'bang':
            x = kiem_gon(md)
            if x:
                gon[md['khoa']] = (md, x)
    if gon:
        goc = tep_goc(du_lieu, so_voi)
        diem_goc = {}
        if goc is not None:
            for m in doc_du_lieu(goc, hang)[1]:
                diem_goc[chuan_ten(m['khoa'])] = sorted((p[0], p[1]) for p in m['diem'])
        co_san, moi = [], []
        for khoa, (md, ds_gon) in gon.items():
            cur = sorted((p[0], p[1]) for p in md['diem'])
            g = diem_goc.get(chuan_ten(khoa))
            giong = g is not None and len(g) == len(cur) and all(
                abs(a[0] - b[0]) < 1e-9 and abs(a[1] - b[1]) < 1e-9 for a, b in zip(g, cur))
            (co_san if giong else moi).extend(f'Model {khoa}: {x}' for x in ds_gon)
        ten_goc = goc.name if goc is not None else '(không có số liệu gốc để so)'
        canh_bao += [f'gợn có sẵn trong số liệu đã duyệt — không chặn ({ten_goc}): {x}' for x in co_san]
        if moi:
            goi_y = 'Gợi ý: chỉnh điểm kề (số nguyên, ghi cũ→mới trong DU_LIEU) hoặc hỏi anh.'
            nguon = (f'điểm đã đổi so với số liệu gốc {goc}' if goc is not None
                     else 'không có số liệu gốc DU_LIEU/<cùng tên> để so (dùng --so-voi)')
            if not chap_nhan_gon:
                raise LoiDuLieu(f'DỮ LIỆU TẠO GỢN ({nguon}) — ' + ' | '.join(moi) +
                                f'. {goi_y} Vẽ bất chấp: --chap-nhan-gon.')
            canh_bao += [f'GỢN (vẽ bất chấp do --chap-nhan-gon; {nguon}): {x}. {goi_y}' for x in moi]
    W = float(rong or S.rong_md)
    ph = Phong(S.phong, cho_phep_gan_giong)
    if ph.gan_giong:
        canh_bao.insert(0, f'gần giống: {ph.ho_thuc[False]} / {ph.ho_thuc[True]} thay cho "{S.phong}" '
                           f'(--cho-phep-phong-gan-giong) — KHÔNG đúng hồ sơ hãng.')
    dk = khoi_dieu_kien(S, tt, canh_bao)
    if kieu == 'bang':
        nhom = [ds]
        if cao:
            B = dung_bang(S, ph, ds, W, [], bo_cot_a)
            cho = float(cao) - (cao_tieu_de(S, ph) + 4 if tieu_de else 0) - 2 * B['cao_hang']
            sh = max(1, int(cho // B['cao_hang'])) if B['huong'] == 'dong' else len(ds)
            if B['huong'] == 'cot' and B['cao'] > float(cao):
                raise LoiBoCuc(f'Bảng cao {B["cao"]:.0f} pt > {float(cao):.0f} pt — cần tách bớt cột Q.')
            nhom = chia_nhom(ds, sh)
    else:
        nhom = chia_nhom(ds, S.toi_da_bd)
    if len(nhom) > 1:
        canh_bao.append(f'{len(ds)} model > {S.toi_da_bd if kieu != "bang" else "sức chứa"} mỗi '
                        f'{"biểu đồ" if kieu != "bang" else "trang"} — tách {len(nhom)} trang: '
                        + ' + '.join(str(len(n)) for n in nhom))
    doc = pymupdf.open()
    dung = {False: set(), True: set()}
    bc = dict(ra=str(ra), hang=hang, kieu=kieu, phong=ph.mo_ta(), phong_gan_giong=ph.gan_giong,
              kho=(W, None), trang=[], canh_bao=canh_bao, dieu_kien=dk)
    for so, nh in enumerate(nhom, start=1):
        phan = f'{so}/{len(nhom)}' if len(nhom) > 1 else ''
        cb = []
        B = dung_bang(S, ph, nh, W, cb, bo_cot_a) if kieu in ('bang', 'ca_hai') else None
        h_td = (cao_tieu_de(S, ph) + 6) if tieu_de else 0.0
        if kieu == 'bang':
            H = float(cao) if cao else h_td + B['cao'] + 1
        else:
            H = float(cao or S.cao_md)
        K = Khung(doc, W, H, ph, S, dung)
        y = 0.0
        tt_trang = dict(so=so, model=[md['khoa'] for md in nh])
        mau_so = None
        if kieu in ('bieu_do', 'ca_hai'):
            if tieu_de:
                ve_tieu_de(K, S, 0, y, W, 'PERFORMANCE CURVES', phan)
                y += h_td
            h_bd = H - y - ((h_td + B['cao'] + 12) if kieu == 'ca_hai' else 0)
            kq = ve_bieu_do(K, S, pymupdf.Rect(0, y, W, y + h_bd), nh, dk, cb)
            tt_trang.update(kq)
            mau_so = [mau(kq['mau'][md['khoa']]) for md in nh]
            y += h_bd + 12
        if kieu in ('bang', 'ca_hai'):
            if tieu_de:
                ve_tieu_de(K, S, 0, y, W, 'PERFORMANCE DATA',
                           chuoi_dien(S, nh) + (f'  ·  {phan}' if phan and kieu == 'bang' else ''))
                y += h_td
            if mau_so is None and B['huong'] == 'cot':
                mau_so = mau_duong(S, len(nh), [])
            ve_bang(K, S, B, 0, y, mau_so)
            tt_trang['bang'] = dict(co=B['co'], rong=round(B['rong'], 1), cao=round(B['cao'], 1),
                                    cot_q=[fmt(q) for q in B['cot_q']],
                                    cot=[f'{a} {b}'.strip() for a, b, *_ in B['tt']])
            y += B['cao']
        tt_trang['tp_min'] = round(K.tu_kiem(), 2)
        tt_trang['kho'] = (W, H)
        for c in cb:
            if c not in canh_bao:
                canh_bao.append(c)
        bc['trang'].append(tt_trang)
    doc.subset_fonts()
    sua_tounicode(doc, ph, dung)
    Path(ra).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ra), garbage=4, deflate=True)
    doc.close()
    bc['tp_min'] = min(t['tp_min'] for t in bc['trang'])
    bc['so_trang'] = len(bc['trang'])
    bc['dung_luong'] = os.path.getsize(ra)
    bc['tep_canh_bao'] = ghi_canh_bao(ra, canh_bao, f've_duong_cong.py — {Path(ra).name} — {hang} — ĐẠT, '
                                                    f'{bc["so_trang"]} trang. Phông: {bc["phong"]}')
    return bc


def _gioi_han_duong():
    try:
        ch = json.loads(HO_SO_MAC_DINH.read_text(encoding='utf-8'))['chung']
        n, m = int(ch['so_duong_toi_da_moi_bieu_do']), len(ch['okabe_ito_bo_vang'])
    except (OSError, ValueError, KeyError):
        n, m = 8, 7
    return (f'SỐ ĐƯỜNG TỐI ĐA MỖI BIỂU ĐỒ: AQUARIS {min(n, m)} (bảng màu Okabe–Ito bỏ vàng chỉ có {m} màu); '
            f'PEROTAC, REDPUMP {n}. Nhiều model hơn -> tự tách thành nhiều trang (mỗi trang một biểu đồ + bảng).')


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except (AttributeError, ValueError):
        pass
    ap = argparse.ArgumentParser(
        description='Vẽ biểu đồ Q–H / bảng thông số theo hồ sơ hãng (PDF vector, chữ thật, phông nhúng).',
        epilog=_gioi_han_duong() + '\nPhông tìm theo thứ tự: <thư mục script>/../fonts, <gốc>/vendor/fonts, hệ thống. '
               'Mã thoát: 0 đạt; 2 dữ liệu sai/gợn; 3 không vừa, thiếu phông, nhãn không đặt được; 1 lỗi khác. '
               'Mọi cảnh báo ghi thêm vào <ra>.canh_bao.txt.\n'
               'Gợn: dãy (Q,H) giống hệt số liệu gốc đã duyệt (--so-voi, mặc định DU_LIEU/<cùng tên tệp>; tệp vào nằm '
               'trong DU_LIEU/ thì so với chính nó) -> chỉ cảnh báo "gợn có sẵn trong số liệu đã duyệt — không chặn"; '
               'điểm đã đổi hoặc không có bản gốc mà tạo gợn -> mã 2 (trừ khi --chap-nhan-gon).',
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--du-lieu', required=True, help='DU_LIEU/<NHÓM>.xlsx')
    ap.add_argument('--hang', required=True, help='AQUARIS (tối đa 7 đường/biểu đồ) | PEROTAC | REDPUMP (tối đa 8)')
    ap.add_argument('--ra', required=True, help='tệp PDF xuất ra')
    ap.add_argument('--model', default='', help='danh sách model, cách nhau dấu phẩy (mặc định: tất cả)')
    ap.add_argument('--kieu', default='ca_hai', choices=['bieu_do', 'bang', 'ca_hai'])
    ap.add_argument('--rong-pt', type=float, default=None)
    ap.add_argument('--cao-pt', type=float, default=None)
    ap.add_argument('--ho-so', default=None, help='mặc định rules/HO_SO_HANG.json cạnh tools/')
    ap.add_argument('--khong-tieu-de', action='store_true', help='không vẽ dòng tiêu đề khối')
    ap.add_argument('--cho-phep-phong-gan-giong', action='store_true',
                    help='thiếu phông đúng hồ sơ hãng thì dùng phông gần giống (in cảnh báo) thay vì dừng mã 3')
    ap.add_argument('--chap-nhan-gon', action='store_true',
                    help='vẫn vẽ khi dữ liệu tạo gợn (in cảnh báo) thay vì dừng mã 2')
    ap.add_argument('--so-voi', default=None,
                    help='xlsx số liệu gốc đã duyệt để so khi gặp gợn (mặc định DU_LIEU/<cùng tên tệp> của kho)')
    ap.add_argument('--bo-cot-a', action='store_true',
                    help='(cũ, không cần nữa) nguồn không có A thì công cụ tự bỏ cột A theo QĐ 152; giữ cờ để lệnh cũ vẫn chạy')
    a = ap.parse_args(argv)
    models = [m.strip() for m in a.model.split(',') if m.strip()] or None
    try:
        bc = ve(a.du_lieu, a.hang, a.ra, models, a.kieu, a.rong_pt, a.cao_pt, a.ho_so, not a.khong_tieu_de,
                a.cho_phep_phong_gan_giong, a.chap_nhan_gon, a.bo_cot_a, a.so_voi)
    except LoiVe as e:
        nhan = {2: 'LỖI DỮ LIỆU', 3: 'KHÔNG VỪA / THIẾU / CẦN TÁCH'}.get(e.ma, 'LỖI')
        print(f'{nhan}: {e}')
        ghi_canh_bao(a.ra, [f'{nhan}: {e}'], f've_duong_cong.py — {Path(a.ra).name} — {a.hang} — KHÔNG XUẤT (mã {e.ma})')
        return e.ma
    n_model = sum(len(t['model']) for t in bc['trang'])
    print(f'ĐẠT: {bc["ra"]} — {bc["hang"]}, {bc["so_trang"]} trang, {n_model} model, '
          f'{bc["dung_luong"] / 1024:.0f} KB, tương phản chữ nhỏ nhất {bc["tp_min"]:.2f}:1')
    print(f'Phông: {bc["phong"]}')
    for t in bc['trang']:
        kho = t['kho']
        print(f'  Trang {t["so"]} ({kho[0]:.0f}×{kho[1]:.0f} pt): ' + ', '.join(t['model']))
    for c in bc['canh_bao']:
        print(f'CẢNH BÁO: {c}')
    if bc['canh_bao']:
        print(f'(Đã ghi {len(bc["canh_bao"])} cảnh báo vào {bc["tep_canh_bao"]})')
    return 0


if __name__ == '__main__':
    sys.exit(main())
