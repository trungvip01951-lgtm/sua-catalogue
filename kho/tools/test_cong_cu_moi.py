#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_cong_cu_moi.py - bài kiểm nhỏ cho kiem_catalogue.py, nhan_ban_chinh.py.

Tự tạo PDF giả bằng pymupdf trong thư mục tạm NGOÀI kho:
  - 1 PDF đạt (A4, có "50 Hz", biểu đồ Q–H 3 đường + bảng 3 model khớp nhau, ảnh 540 ppi);
  - 1 PDF lỗi cài sẵn: chữ "warranty", ảnh 72 ppi, khổ lẻ 600×800 pt, đường có móc cuối.
Rồi kiểm: kiem_catalogue bắt đúng (và không báo nhầm trên PDF đạt), --du-lieu so đường ≤5%, --bo3,
nhan_ban_chinh --thu (không ghi gì; PDF lỗi mục 1 phải bị DỪNG).
Dùng: python -B -X utf8 tools/test_cong_cu_moi.py      Mã thoát 0 = tất cả ĐẠT.
"""
from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import pymupdf
from scipy.interpolate import PchipInterpolator

TOOLS = Path(__file__).resolve().parent
KHO = TOOLS.parent
PY = [sys.executable, '-B', '-X', 'utf8']
DU_LIEU = {  # model: (kW, [(Q m³/h, H m)])
    'TN50/05': (0.37, [(0, 18), (2, 16), (4, 13), (6, 8)]),
    'TN50/07': (0.55, [(0, 24), (2, 22), (4, 19), (6, 14), (7, 10)]),
    'TN50/10': (0.75, [(0, 29), (2, 27), (4, 24), (6, 19), (8, 12)]),
}
QCOT = [0, 2, 4, 6, 8]
KQ = []


def ghi(ten, dat, chi_tiet=''):
    KQ.append(dat)
    print(f'  [{"ĐẠT" if dat else "CHƯA ĐẠT"}] {ten}' + (f' — {chi_tiet}' if chi_tiet else ''))


def chu(pg, x, y, s, size=10, giua=False, phai=False):
    w = pymupdf.get_text_length(s, fontname='helv', fontsize=size)
    if giua:
        x -= w / 2
    if phai:
        x -= w
    pg.insert_text((x, y), s, fontname='helv', fontsize=size, color=(0, 0, 0))


def tao_pdf(path, loi=False):
    doc = pymupdf.open()
    W, H = (600, 800) if loi else (595, 842)
    pg = doc.new_page(width=W, height=H)
    chu(pg, 40, 50, 'TN SERIES CENTRIFUGAL PUMP', 18)
    chu(pg, 40, 80, '50 Hz  -  2-pole motor  -  3~ 380-415 V 50 Hz', 10)  # phông helv không có dấu gạch ngang dài
    than = ['Single-stage centrifugal pump for clean water supply and pressure boosting duties.',
            'Pump body and impeller made of stainless steel for long service life in use.',
            'Motor with insulation class F and protection degree IP55 for continuous duty.',
            'Typical values shown below are measured with clean water at normal temperature.']
    if loi:
        than[2] = 'Five-year warranty on all models of this series for every customer.'
    for k, s in enumerate(than):
        chu(pg, 40, 110 + 18 * k, s, 10)
    # ảnh
    px, rect = ((200, 133), pymupdf.Rect(380, 95, 580, 228)) if loi else ((1200, 800), pymupdf.Rect(400, 100, 560, 206.7))
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, *px), False)
    pix.set_rect(pix.irect, (40, 90, 140))
    pix.set_rect(pymupdf.IRect(px[0] // 4, px[1] // 4, px[0] * 3 // 4, px[1] * 3 // 4), (220, 220, 220))
    pg.insert_image(rect, pixmap=pix)
    # biểu đồ Q–H: Q 0..8 m³/h ↔ x 100..500; H 0..30 m ↔ y 480..240
    X = lambda q: 100 + q * 50  # noqa
    Y = lambda h: 480 - h * 8  # noqa
    sh = pg.new_shape()
    for h in range(0, 31, 5):
        sh.draw_line((100, Y(h)), (500, Y(h)))
    for q in range(0, 9):
        sh.draw_line((X(q), 240), (X(q), 480))
    sh.finish(color=(0.75, 0.75, 0.75), width=0.5)
    sh.commit()
    for h in range(0, 31, 5):
        chu(pg, 94, Y(h) + 2.8, str(h), 8, phai=True)
    for q in range(0, 9):
        chu(pg, X(q), 494, str(q), 8, giua=True)
    chu(pg, 30, 494, 'Q (m³/h)', 8)
    chu(pg, 70, 228, 'H (m)', 8)
    mau = [(0, 0.45, 0.70), (0.84, 0.37, 0), (0, 0.62, 0.45)]
    for k, (ten, (kw, pts)) in enumerate(DU_LIEU.items()):
        q = np.array([p[0] for p in pts], float)
        hh = np.array([p[1] for p in pts], float)
        f = PchipInterpolator(q, hh)
        qq = np.linspace(q[0], q[-1], 40)
        P = [(X(a), Y(b)) for a, b in zip(qq, f(qq))]
        if loi and k == 0:  # đoạn gần phẳng rồi dốc gắt
            P = [(X(a), Y(18 - 0.05 * a)) for a in np.linspace(0, 4, 20)] + [(X(a), Y(17.8 - 4.9 * (a - 4))) for a in np.linspace(4.1, 6, 10)]
        if loi and k == 1:  # gợn: nhiều điểm uốn (vẫn H giảm khi Q tăng)
            qq2 = np.linspace(0, 7, 60)
            P = [(X(a), Y(24 - 2 * a - 1.2 * np.sin(1.5 * a))) for a in qq2]
        if loi and k == 2:
            P += [(X(8.15), Y(12.8)), (X(8.3), Y(14.5))]  # móc cuối: H tăng lại
        sh = pg.new_shape()
        sh.draw_polyline(P)
        sh.finish(color=mau[k], width=1.6, closePath=False)
        sh.commit()
        chu(pg, P[-1][0] + 4, P[-1][1] - 3, ten, 7.5)
    # bảng
    chu(pg, 60, 540, 'MODEL', 8)
    chu(pg, 165, 540, 'kW', 8)
    chu(pg, 215, 540, 'm³/h', 8)
    xc = [270, 320, 370, 420, 470]
    for x, q in zip(xc, QCOT):
        chu(pg, x, 540, str(q), 8, giua=True)
    chu(pg, 215, 552, 'l/min', 8)
    for x, q in zip(xc, QCOT):
        chu(pg, x, 552, f'{q * 16.667:.0f}', 8, giua=True)
    for r, (ten, (kw, pts)) in enumerate(DU_LIEU.items()):
        y = 570 + 14 * r
        chu(pg, 60, y, ten, 8)
        chu(pg, 165, y, f'{kw:g}', 8)
        d = dict(pts)
        for x, q in zip(xc, QCOT):
            chu(pg, x, y, f'{d[q]:g}' if q in d else '–', 8, giua=True)
    cuoi = ['Curves at 50 Hz, 2 poles, clean water at 20 °C, kinematic viscosity 1 mm²/s.',
            'Typical values, not guaranteed. Performance tolerance according to the standard.',
            'Suction and discharge ports are threaded, flange adapters are available on request.',
            'All dimensions are given in millimetres and subject to change without notice.',
            'Electrical connection must be carried out by qualified personnel only, always.',
            'Install the pump in a dry and ventilated place, protected from frost and rain.',
            'Check the direction of rotation before the first start of the three-phase motor.',
            'Never run the pump dry; fill the pump body with water before starting the motor.',
            'Keep the suction pipe as short as possible and avoid air pockets in the pipe line.',
            'For other voltages and special versions please contact our technical department.',
            'TN SERIES  /  TECHNICAL DATA  /  PAGE 01']
    for k, s in enumerate(cuoi):
        chu(pg, 40, (640 if not loi else 620) + 15 * k, s, 9 if k < 10 else 9)
    doc.save(str(path), garbage=4, deflate=True)
    doc.close()


def tao_pdf_wc(path, h_cuoi):
    # bảng WC: 2 cỡ họng; H tại Qmax theo QĐ 19 = nấc công suất (1, 2, 3 …) nếu h_cuoi đúng
    doc = pymupdf.open()
    pg = doc.new_page(width=842, height=595)
    chu(pg, 40, 50, 'WC SERIES CUTTER PUMP  -  50 Hz', 16)
    chu(pg, 60, 100, 'MODEL', 8)
    chu(pg, 200, 100, 'kW', 8)
    chu(pg, 250, 100, 'm³/h', 8)
    xc = [300, 350, 400, 450, 500]
    for x, q in zip(xc, [0, 10, 20, 30, 40]):
        chu(pg, x, 100, str(q), 8, giua=True)
    hang = [('50WC15-15-1.5', 1.5, [18, 16, 12, 6]), ('50WC9-22-2.2', 2.2, [24, 22, 18, 10]),
            ('50WC15-25-3', 3, [30, 27, 22, 14]), ('65WC25-10-1.5', 1.5, [14, 12, 9, 5]),
            ('65WC28-15-2.2', 2.2, [20, 18, 15, 10])]
    for r, (ten, kw, hs) in enumerate(hang):
        y = 120 + 14 * r
        chu(pg, 60, y, ten, 8)
        chu(pg, 200, y, f'{kw:g}', 8)
        hs = hs + [h_cuoi[r]]
        for k, x in enumerate(xc):
            chu(pg, x, y, f'{hs[k]:g}' if k < len(hs) else '–', 8, giua=True)
    doc.save(str(path))
    doc.close()


def tao_du_lieu_vdc(path):
    # 5 model giả, dạng DU_LIEU của ve_duong_cong.py (THONG_TIN / MODEL / QH)
    import openpyxl
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'THONG_TIN'
    s.append(['khoa', 'gia_tri', 'nguon'])
    s.append(['so_cuc', 2, 'giả'])
    s.append(['toc_do_vg_ph', 2900, 'giả'])
    m = wb.create_sheet('MODEL')
    m.append(['hang', 'model_1pha', 'model_3pha', 'p2_kw', 'p2_hp', 'dong_A_1pha', 'dong_A_3pha', 'dien_ap_1pha',
              'dien_ap_3pha', 'hong_hut', 'hong_xa', 'thong_so_khac', 'nguon', 'nghi_ngo'])
    q = wb.create_sheet('QH')
    q.append(['hang', 'model', 'q_m3h', 'h_m', 'nguon'])
    luoi = [0, 4, 8, 12, 16, 20]
    for k in range(5):
        qmax = [8, 8, 12, 16, 20][k]
        h0 = 12 + 6.5 * k
        m3, m1 = f'TK {50 + 5 * (k // 2)}-{k + 1}', (f'TKM {50 + 5 * (k // 2)}-{k + 1}' if k < 2 else None)
        m.append([None, m1, m3, [0.37, 0.55, 0.75, 1.1, 1.5][k], [0.5, 0.75, 1, 1.5, 2][k], (2.6 + k if k < 2 else None),
                  round(1 + 0.6 * k, 1), '1~ 220–240 V 50 Hz' if m1 else None, '3~ 380–415 V 50 Hz', 50, 32, None, 'giả', None])
        for qq in luoi:
            if qq <= qmax:
                q.append([None, m3, qq, round(h0 * (1 - 0.35 * (qq / qmax) ** 2), 1), 'giả'])
    wb.save(str(path))


def tao_xlsx(path, he_so=1.0):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'QH'
    ws.append(['hang', 'model', 'q_m3h', 'h_m', 'nguon'])
    for ten, (kw, pts) in DU_LIEU.items():
        for q, h in pts:
            ws.append(['', ten, q, h * (he_so if ten == 'TN50/10' else 1.0), 'thử'])
    wb.save(str(path))


def chay(args, ten_tep_json=None):
    p = subprocess.run(PY + args, capture_output=True, text=True, encoding='utf-8', cwd=str(KHO))
    js = json.loads(Path(ten_tep_json).read_text(encoding='utf-8')) if ten_tep_json and Path(ten_tep_json).exists() else None
    return p.returncode, p.stdout + p.stderr, js


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest() if Path(p).exists() else None


def main():
    tmp = Path(tempfile.mkdtemp(prefix='test_cong_cu_'))
    print(f'== BÀI KIỂM CÔNG CỤ MỚI (thư mục tạm: {tmp})')
    tot, xau = tmp / 'AQUARIS - Bơm thử nghiệm TN.pdf', tmp / 'loi' / 'AQUARIS - Bơm thử nghiệm TN.pdf'
    xau.parent.mkdir()
    tao_pdf(tot)
    tao_pdf(xau, loi=True)
    # 1) PDF đạt: mục 1–8 phải ĐẠT
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(tot), '--json', str(tmp / 'tot.json')], tmp / 'tot.json')
    if not js:
        ghi('kiem_catalogue chạy trên PDF đạt', False, out[-400:])
        return 1
    m = js['tep'][0]['muc']
    for k in range(1, 9):
        ghi(f'PDF đạt — mục {k} {m[str(k)]["ten"]}: {m[str(k)]["trang_thai"]}', m[str(k)]['trang_thai'] == 'ĐẠT',
            '; '.join((m[str(k)]['loi'] + m[str(k)]['chua_the_xac_minh'])[:2])[:200])
    ghi('PDF đạt — không áp QĐ 19 cho bơm không phải WC', not any('QĐ 19' in g for g in m['3']['ghi_chu'] + m['3']['loi']))
    ghi('PDF đạt — mục 3 đọc đúng 3 model / 3 đường / 3 nhãn', any('3 model / 3 đường / 3 nhãn' in g for g in m['3']['ghi_chu']),
        ' | '.join(m['3']['ghi_chu'])[:160])
    # 2) PDF lỗi: phải bắt đúng từng lỗi cài sẵn
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(xau), '--json', str(tmp / 'xau.json')], tmp / 'xau.json')
    m = js['tep'][0]['muc']
    bat = [(1, 'chữ "warranty"', lambda l: 'warranty' in l),
           (6, 'ảnh 72 ppi', lambda l: '= 72 ppi' in l),
           (8, 'khổ lẻ 600×800', lambda l: 'không phải A4' in l and '600.0×800.0' in l),
           (4, 'đường có móc', lambda l: 'móc' in l),
           (4, 'đoạn phẳng rồi dốc gắt', lambda l: 'không cong đều' in l and 'dốc gắt' in l),
           (4, 'nhiều điểm uốn', lambda l: 'không cong đều' in l and 'điểm uốn' in l)]
    for k, ten, f in bat:
        l = [x for x in m[str(k)]['loi'] if f(x)]
        ghi(f'PDF lỗi — mục {k} bắt {ten}', m[str(k)]['trang_thai'] == 'CHƯA ĐẠT' and bool(l), (l or m[str(k)]['loi'] or ['-'])[0][:160])
    ghi('PDF lỗi — mã thoát 1 (có CHƯA ĐẠT)', rc == 1, f'rc={rc}')
    # 3) --du-lieu: khớp ≤5% thì ĐẠT; H lệch +20% thì CHƯA ĐẠT
    tao_xlsx(tmp / 'TN.xlsx')
    tao_xlsx(tmp / 'TN_lech.xlsx', 1.2)
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(tot), '--muc', '4', '--du-lieu', str(tmp / 'TN.xlsx'),
                        '--json', str(tmp / 'dl.json')], tmp / 'dl.json')
    m4 = js['tep'][0]['muc']['4']
    ghi('--du-lieu khớp: mục 4 ĐẠT và đã so 3 model', m4['trang_thai'] == 'ĐẠT' and any('so 3 model' in g for g in m4['ghi_chu']),
        '; '.join(m4['loi'] + m4['chua_the_xac_minh'])[:160])
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(tot), '--muc', '4', '--du-lieu', str(tmp / 'TN_lech.xlsx'),
                        '--json', str(tmp / 'dl2.json')], tmp / 'dl2.json')
    m4 = js['tep'][0]['muc']['4']
    ghi('--du-lieu lệch 20%: mục 4 CHƯA ĐẠT ở TN50/10', m4['trang_thai'] == 'CHƯA ĐẠT' and any('TN50/10' in l and '> 5%' in l for l in m4['loi']),
        (m4['loi'] or ['-'])[0][:160])
    # 3b) QĐ 19 WC: đúng nấc → ĐẠT; sai một mã → CHƯA ĐẠT
    tao_pdf_wc(tmp / 'PEROTAC - Bơm cắt rác WC thử.pdf', [1, 2, 3, 1, 2])
    (tmp / 'sai').mkdir()
    tao_pdf_wc(tmp / 'sai' / 'PEROTAC - Bơm cắt rác WC thử.pdf', [1, 2, 4, 1, 2])
    for ten, duong, ky in (('QĐ 19 WC đúng nấc (1,2,3 | 1,2): mục 3 ĐẠT', tmp / 'PEROTAC - Bơm cắt rác WC thử.pdf', 'ĐẠT'),
                           ('QĐ 19 WC sai (50WC15-25-3 có 4 m, phải 3 m): mục 3 CHƯA ĐẠT', tmp / 'sai' / 'PEROTAC - Bơm cắt rác WC thử.pdf', 'CHƯA ĐẠT')):
        rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(duong), '--muc', '3', '--json', str(tmp / 'wc.json')], tmp / 'wc.json')
        m3 = js['tep'][0]['muc']['3']
        ok = m3['trang_thai'] == ky and any('QĐ 19' in g for g in m3['ghi_chu'])
        if ky == 'CHƯA ĐẠT':
            ok = ok and any('50WC15-25-3' in l and 'phải = 3 m' in l for l in m3['loi'])
        ghi(ten, ok, '; '.join(m3['loi'] + [g for g in m3['ghi_chu'] if 'QĐ 19' in g])[:170])
    # 3c) PDF do tools/ve_duong_cong.py vẽ (5 model, 3 hãng) → kiem --du-lieu ĐẠT mục 2, 3, 4
    vdc = TOOLS / 've_duong_cong.py'
    if vdc.exists():
        tao_du_lieu_vdc(tmp / 'TK.xlsx')
        for h in ('AQUARIS', 'PEROTAC', 'REDPUMP'):
            pdf = tmp / 'vdc' / f'{h} - Bơm thử TK.pdf'
            pdf.parent.mkdir(exist_ok=True)
            p = subprocess.run(PY + [str(vdc), '--du-lieu', str(tmp / 'TK.xlsx'), '--hang', h, '--ra', str(pdf), '--kieu', 'ca_hai'],
                               capture_output=True, text=True, encoding='utf-8', cwd=str(KHO))
            if p.returncode != 0:
                ghi(f've_duong_cong vẽ {h}', False, (p.stdout + p.stderr)[-200:])
                continue
            rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(pdf), '--muc', '2,3,4', '--du-lieu', str(tmp / 'TK.xlsx'),
                                '--json', str(tmp / 'vdc.json')], tmp / 'vdc.json')
            mm = js['tep'][0]['muc']
            ok = all(mm[k]['trang_thai'] == 'ĐẠT' for k in ('2', '3', '4')) and \
                any('so 5 model' in g for g in mm['4']['ghi_chu'])
            ghi(f'PDF do ve_duong_cong vẽ ({h}): mục 2/3/4 = {mm["2"]["trang_thai"]}/{mm["3"]["trang_thai"]}/{mm["4"]["trang_thai"]}',
                ok, '; '.join(mm['2']['loi'] + mm['3']['loi'] + mm['4']['loi'] + mm['4']['chua_the_xac_minh']
                              + [g for g in mm['4']['ghi_chu'] if 'so 5' in g])[:170])
    else:
        ghi('có tools/ve_duong_cong.py để kiểm chéo', False, 'không thấy tệp')
    # 3d) PEROTAC CM đã duyệt so DU_LIEU/CM.xlsx: không model nào lệch > 5%
    cm = KHO / 'CATALOG SUA' / 'CATALOGUE DA DUYET' / 'PEROTAC' / 'PEROTAC - Bơm trục ngang CM.pdf'
    if cm.exists() and (KHO / 'DU_LIEU' / 'CM.xlsx').exists():
        rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(cm), '--muc', '4', '--du-lieu', str(KHO / 'DU_LIEU' / 'CM.xlsx'),
                            '--json', str(tmp / 'cm.json')], tmp / 'cm.json')
        m4 = js['tep'][0]['muc']['4']
        lech = [l for l in m4['loi'] if '> 5%' in l or 'không có đường nào' in l]
        g = [x for x in m4['ghi_chu'] if 'lệch lớn nhất' in x]
        ghi('PEROTAC CM đã duyệt so DU_LIEU/CM.xlsx ≤5%', not lech and not m4['chua_the_xac_minh'] and bool(g),
            (g or ['-'])[0][:120] + (f'; LỆCH: {lech[:3]}' if lech else ''))
    # 3e) --goc: chỉ tính lỗi mới so với bản gốc
    va_ = tmp / 'va' / 'AQUARIS - Bơm thử nghiệm TN.pdf'
    va_.parent.mkdir()
    d = pymupdf.open(str(tot))
    d[0].insert_text((40, 236), 'Two-year warranty.', fontname='helv', fontsize=9)
    d.save(str(va_))
    d.close()
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(va_), '--goc', str(tot), '--muc', '1', '--json', str(tmp / 'g1.json')], tmp / 'g1.json')
    m1 = js['tep'][0]['muc']['1']
    ghi('--goc: chữ "warranty" mới thêm → mục 1 CHƯA ĐẠT', m1['trang_thai'] == 'CHƯA ĐẠT' and any('warranty' in l for l in m1['loi']))
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(va_), '--goc', str(tot), '--trang', '2', '--muc', '1',
                        '--json', str(tmp / 'g2.json')], tmp / 'g2.json')
    m1 = js['tep'][0]['muc']['1']
    ghi('--goc --trang 2: lỗi ở tr.1 (ngoài trang vá) không tính', m1['trang_thai'] == 'ĐẠT' and
        any('ngoài trang' in l for l in m1['loi_co_tu_ban_goc']))
    xau2 = tmp / 'va2' / xau.name
    xau2.parent.mkdir()
    d = pymupdf.open(str(xau))
    d.set_metadata({'title': 'bản vá'})
    d.save(str(xau2))
    d.close()
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), str(xau2), '--goc', str(xau), '--muc', '1,4,6,8',
                        '--json', str(tmp / 'g3.json')], tmp / 'g3.json')
    mm = js['tep'][0]['muc']
    ghi('--goc: lỗi có sẵn ở bản gốc (warranty, móc, 72 ppi, khổ lẻ) không tính → mục 1/4/6/8 ĐẠT',
        all(mm[k]['trang_thai'] == 'ĐẠT' and mm[k]['loi_co_tu_ban_goc'] for k in ('1', '4', '6', '8')),
        ' / '.join(f'{k}:{mm[k]["trang_thai"]}' for k in ('1', '4', '6', '8')))
    # 4) --bo3: 3 bản giống hệt → mục 9 CHƯA ĐẠT
    b3 = []
    for h in ('AQUARIS', 'PEROTAC', 'REDPUMP'):
        p = tmp / 'bo3' / f'{h} - Bơm thử nghiệm TN.pdf'
        p.parent.mkdir(exist_ok=True)
        shutil.copyfile(tot, p)
        b3.append(str(p))
    rc, out, js = chay([str(TOOLS / 'kiem_catalogue.py'), '--bo3', *b3, '--muc', '9', '--json', str(tmp / 'b3.json')], tmp / 'b3.json')
    m9 = js['tep'][0]['muc']['9']
    ghi('--bo3 ba bản giống hệt: mục 9 CHƯA ĐẠT', m9['trang_thai'] == 'CHƯA ĐẠT' and 'SSIM 1.00' in ' '.join(m9['loi']),
        (m9['loi'] or ['-'])[0][:120])
    # 5) nhan_ban_chinh --thu: không ghi gì; PDF lỗi mục 1 phải dừng
    ds = KHO / 'CATALOG SUA' / 'CATALOGUE DA DUYET' / 'DANH SACH.txt'
    td = KHO / 'CATALOG SUA' / 'BẢN MỚI 03-10-2026 (ChatGPT)' / '_BAO CAO' / 'TIEN DO.md'
    dich = KHO / 'CATALOG SUA' / 'CATALOGUE DA DUYET' / 'AQUARIS' / 'AQUARIS - Bơm thử nghiệm TN.pdf'
    hist = KHO / 'history'
    truoc = (sha(ds), sha(td), dich.exists(), sorted(p.name for p in hist.iterdir()) if hist.is_dir() else [])
    rc, out, _ = chay([str(TOOLS / 'nhan_ban_chinh.py'), '--ung-vien', str(tot), '--dich', 'AQUARIS/AQUARIS - Bơm thử nghiệm TN.pdf',
                       '--ly-do', 'bài kiểm công cụ', '--thu'])
    sau = (sha(ds), sha(td), dich.exists(), sorted(p.name for p in hist.iterdir()) if hist.is_dir() else [])
    ghi('nhan_ban_chinh --thu (PDF đạt): mã 0, không ghi gì', rc == 0 and 'CHẠY THỬ XONG' in out and truoc == sau,
        f'rc={rc}; ' + out.strip().splitlines()[-1][:120])
    rc, out, _ = chay([str(TOOLS / 'nhan_ban_chinh.py'), '--ung-vien', str(xau), '--dich', 'AQUARIS/AQUARIS - Bơm thử nghiệm TN.pdf',
                       '--ly-do', 'bài kiểm công cụ', '--thu'])
    ghi('nhan_ban_chinh --thu (PDF lỗi): DỪNG vì mục 1, mã 3', rc == 3 and 'DỪNG' in out, f'rc={rc}')
    n = sum(KQ)
    print(f'KẾT QUẢ: {n}/{len(KQ)} ĐẠT')
    shutil.rmtree(tmp, ignore_errors=True)
    return 0 if n == len(KQ) else 1


if __name__ == '__main__':
    sys.exit(main())
