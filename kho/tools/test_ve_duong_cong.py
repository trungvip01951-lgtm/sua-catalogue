#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""test_ve_duong_cong.py — kiểm tools/ve_duong_cong.py bằng dữ liệu giả.

Chạy:  python tools/test_ve_duong_cong.py        (mã thoát 0 = đạt hết)
       python tools/test_ve_duong_cong.py --giu   (giữ PDF thử trong thư mục tạm để xem)

Nội dung:
  1. 5 model giả (2 model có bản 1 pha + dòng điện A) -> vẽ biểu đồ + bảng cho AQUARIS, PEROTAC, REDPUMP: đạt.
  2. Bộ sai (H tăng khi Q tăng) -> bị từ chối, thông báo nêu đúng model + điểm; CLI trả mã 2.
  3. 10 model -> tự tách trang (≤ 8 đường/biểu đồ; AQUARIS ≤ 7 màu Okabe–Ito).
  4. Đo trên PDF xuất ra: khung chữ không chồng nhau; chữ không đè nét đường cong; phông nhúng dạng rút gọn
     đúng họ phông hãng; cỡ chữ ≥ ngưỡng; "–" chỉ ở cột Model 1~; chữ trắng trên đầu bảng tối REDPUMP ≥ 4,5:1.
  7. Phông: lấy từ vendor/fonts; thiếu phông hồ sơ -> mã 3 "thiếu phông", --cho-phep-phong-gan-giong -> cảnh báo.
  8. Gợn: ca thật CM 32-160/2.2 (H tại Q 9: 28.5 -> 29.5) bị chặn mã 2, báo model + điểm + gợi ý; --chap-nhan-gon vẽ.
  9. REDPUMP: 7 đường cùng điểm cuối sát nhau (kiểu 4SP3) -> mọi ô tên nằm ở CUỐI PHẢI đường.
 10. Cột A: nguồn không có A -> bỏ cột (QĐ 152); có A nhưng 1 model thiếu -> giữ cột + cảnh báo "THIẾU DÒNG ĐIỆN A" + tệp <ra>.canh_bao.txt.
 11. --help ghi rõ AQUARIS tối đa 7 đường/biểu đồ.
 12. Gợn so với số liệu gốc: giống hệt bản gốc -> chỉ cảnh báo; điểm đã đổi -> mã 2; tệp vào trong DU_LIEU/ -> so với
     chính nó (chỉ cảnh báo); không ghi --so-voi -> tự lấy DU_LIEU/<cùng tên>; số liệu thật DU_LIEU/CM.xlsx (nếu có).
Hàm kiem_pdf() dùng lại được cho PDF thật (vd. mẫu CM).
"""
import contextlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import openpyxl
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ve_duong_cong as V  # noqa: E402

HANG = ('AQUARIS', 'PEROTAC', 'REDPUMP')
COT_MODEL = ['hang', 'model_1pha', 'model_3pha', 'p2_kw', 'p2_hp', 'dong_A_1pha', 'dong_A_3pha', 'dien_ap_1pha',
             'dien_ap_3pha', 'hong_hut', 'hong_xa', 'thong_so_khac', 'nguon', 'nghi_ngo']
Q_LUOI = [0, 4, 8, 12, 16, 20]


def mau_gia(n, h0=12.0, buoc_h=6.5):
    """n model giả: H = h0·(1 − 0,35·(Q/Qmax)²) trên lưới Q chung, Qmax tăng dần."""
    ds = []
    for k in range(n):
        qmax = Q_LUOI[min(len(Q_LUOI) - 1, 2 + k * 3 // max(1, n - 1))] if n > 1 else 20
        h_dau = h0 + buoc_h * k
        diem = [(q, round(h_dau * (1 - 0.35 * (q / qmax) ** 2), 1)) for q in Q_LUOI if q <= qmax]
        ds.append(dict(m3=f'TP {50 + 5 * (k // 2)}-{k + 1}', m1=(f'TPM {50 + 5 * (k // 2)}-{k + 1}' if k < 2 else None),
                       kw=[0.37, 0.55, 0.75, 1.1, 1.5, 2.2, 3, 4, 5.5, 7.5][k % 10],
                       hp=[0.5, 0.75, 1, 1.5, 2, 3, 4, 5.5, 7.5, 10][k % 10],
                       a1=(2.6 + k if k < 2 else None), a3=round(1.0 + 0.6 * k, 1), diem=diem))
    return ds


def tao_xlsx(path, ds, hang=None):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'THONG_TIN'
    s.append(['khoa', 'gia_tri', 'nguon'])
    s.append(['nhom', 'TEST', 'dữ liệu giả'])
    s.append(['so_cuc', 2, 'dữ liệu giả'])
    s.append(['toc_do_vg_ph', 2900, 'dữ liệu giả'])
    m = wb.create_sheet('MODEL')
    m.append(COT_MODEL)
    q = wb.create_sheet('QH')
    q.append(['hang', 'model', 'q_m3h', 'h_m', 'nguon'])
    for md in ds:
        m.append([hang, md['m1'], md['m3'], md['kw'], md['hp'], md['a1'], md['a3'],
                  '1~ 220–240 V 50 Hz' if md['m1'] else None, '3~ 380–415 V 50 Hz', 50, 32, None, 'giả', None])
        for qq, hh in md['diem']:
            q.append([hang, md['m3'] or md['m1'], qq, hh, 'giả'])
    wb.save(path)
    return path


def md_tu_diem(ten, diem, kw=1.1, hp=1.5, a3=2.5, m1=None, a1=None):
    return dict(m3=ten, m1=m1, kw=kw, hp=hp, a1=a1, a3=a3, diem=diem)


def ca_cm32_160_22(h_q9):
    """Ca thật DU_LIEU/CM.xlsx (PEROTAC, CM 32-160/2.2) + 2 model kề trơn; h_q9 = H tại Q 9 m³/h."""
    return [md_tu_diem('CM 32-160/1.5', [(0, 25.4), (6, 23.7), (9, 22.5), (15, 18.5), (18, 15.8)], 1.5, 2),
            md_tu_diem('CM 32-160/2.2', [(0, 31), (6, 29.6), (9, h_q9), (15, 24.5), (18, 22), (24, 15)], 2.2, 3),
            md_tu_diem('CM 32-160/3', [(0, 35), (6, 34.3), (9, 32.5), (15, 28), (18, 25.5), (24, 19), (27, 15)], 3, 4)]


def ca_4sp3():
    """7 đường cùng Qmax, điểm cuối sát nhau (kiểu REDPUMP 4SP3-6…25)."""
    ds = []
    for k, (h0, h1) in enumerate([(37, 15), (56, 23), (74, 31), (93, 39), (112, 46), (136, 57), (155, 65)]):
        diem = [(round(0.6 * j, 1), float(round(h0 - (h0 - h1) * (0.6 * j / 4.2) ** 2))) for j in range(8)]
        ds.append(md_tu_diem(f'4SP3-{[6, 9, 12, 15, 18, 22, 25][k]}', diem, [0.37, 0.55, 0.75, 0.92, 1.1, 1.3, 1.5][k],
                             [0.5, 0.75, 1, 1.25, 1.5, 1.75, 2][k]))
    return ds


# ----------------------------------------------------------------------------- đo trên PDF
def do_pdf(path):
    """Trả về danh sách trang: (page, spans, điểm-dày của nét đường cong ≥ 1,5 pt, bề rộng nét)."""
    doc = pymupdf.open(path)
    ra = []
    for page in doc:
        spans = []
        for b in page.get_text('dict')['blocks']:
            for ln in b.get('lines', []):
                for sp in ln['spans']:
                    if sp['text'].strip():
                        spans.append(dict(t=sp['text'], r=pymupdf.Rect(sp['bbox']), co=sp['size'],
                                          font=sp['font'], mau=sp['color']))
        diem, w_net = [], 0.0
        for d in page.get_drawings():
            w = d.get('width') or 0.0
            if d.get('color') is None or w < 1.5 - 1e-6:
                continue
            w_net = max(w_net, w)
            for it in d['items']:
                if it[0] == 'l':
                    a, b = it[1], it[2]
                    n = max(2, int(np.hypot(b.x - a.x, b.y - a.y) / 0.25) + 1)
                    diem.append(np.column_stack([np.linspace(a.x, b.x, n), np.linspace(a.y, b.y, n)]))
        ra.append((page, spans, np.vstack(diem) if diem else np.zeros((0, 2)), w_net))
    return doc, ra


def kiem_pdf(path, hang, ho_so=None):
    """Kiểm một PDF do ve_duong_cong.py xuất. Trả về danh sách lỗi (rỗng = đạt)."""
    HS = json.loads(Path(ho_so or V.HO_SO_MAC_DINH).read_text(encoding='utf-8'))
    co_min = float(HS['chung']['co_chu_pt']['bang_min'])
    ho = HS['hang'][hang]['phong'][0].replace(' ', '')
    loi = []
    doc, trang = do_pdf(path)
    for so, (page, spans, P, w_net) in enumerate(trang, start=1):
        # 1. khung chữ không chồng nhau
        for i in range(len(spans)):
            for j in range(i + 1, len(spans)):
                a, b = spans[i]['r'], spans[j]['r']
                if min(a.x1, b.x1) - max(a.x0, b.x0) > 0.3 and min(a.y1, b.y1) - max(a.y0, b.y0) > 0.3:
                    loi.append(f'tr.{so}: chữ chồng "{spans[i]["t"]}" / "{spans[j]["t"]}"')
        # 2. chữ không đè nét đường cong (nới khung bằng nửa bề rộng nét)
        e = w_net / 2.0
        for sp in spans:
            r = sp['r']
            if len(P) and np.any((P[:, 0] > r.x0 - e) & (P[:, 0] < r.x1 + e) & (P[:, 1] > r.y0 - e) & (P[:, 1] < r.y1 + e)):
                loi.append(f'tr.{so}: chữ "{sp["t"]}" đè nét đường cong')
        # 3. cỡ chữ
        for sp in spans:
            if sp['co'] < co_min - 1e-6:
                loi.append(f'tr.{so}: chữ "{sp["t"]}" cỡ {sp["co"]:.1f} pt < {co_min:g} pt')
        # 4. phông nhúng, rút gọn, đúng họ
        for f in page.get_fonts(full=True):
            xref, ext, basefont = f[0], f[1], f[3]
            if ext in ('n/a', '') or not doc.extract_font(xref)[3]:
                loi.append(f'tr.{so}: phông {basefont} không nhúng')
            if '+' not in basefont:
                loi.append(f'tr.{so}: phông {basefont} không rút gọn')
            if ho.lower() not in basefont.replace(' ', '').lower():
                loi.append(f'tr.{so}: phông {basefont} khác hồ sơ hãng ({ho})')
        # 5. "–" chỉ ở cột / hàng Model 1~
        gach = [sp for sp in spans if sp['t'].strip() == '–']
        if gach:
            if HS['hang'][hang]['bang']['huong'] == 'model_theo_cot':
                mot = [sp for sp in spans if sp['t'].strip() == '1~']
                ok = mot and all(abs((g['r'].y0 + g['r'].y1) / 2 - (mot[0]['r'].y0 + mot[0]['r'].y1) / 2) < 1 for g in gach)
            else:
                dau = [sp for sp in spans if sp['t'].strip() == 'Model 1~']
                ok = dau and all(abs(g['r'].x0 - dau[0]['r'].x0) < 1 for g in gach)
            if not ok:
                loi.append(f'tr.{so}: có "–" ngoài cột Model 1~')
    doc.close()
    return loi


# ----------------------------------------------------------------------------- ca kiểm
class KiemVeDuongCong(unittest.TestCase):
    tmp = None

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix='test_ve_'))
        cls.ds5 = mau_gia(5)
        cls.x5 = tao_xlsx(cls.tmp / 'gia5.xlsx', cls.ds5)
        cls.x10 = tao_xlsx(cls.tmp / 'gia10.xlsx', mau_gia(10, h0=10, buoc_h=4.2))
        sai = mau_gia(3)
        sai[1]['diem'] = [(0, 20.0), (4, 21.5), (8, 18.0)]          # H tăng từ Q 0 -> 4
        cls.x_sai = tao_xlsx(cls.tmp / 'sai.xlsx', sai)
        cls.HS = json.loads(V.HO_SO_MAC_DINH.read_text(encoding='utf-8'))

    @classmethod
    def tearDownClass(cls):
        if '--giu' in sys.argv:
            print(f'\nGiữ tệp thử ở: {cls.tmp}')
        else:
            shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_1_ve_dat_3_hang(self):
        """5 model giả -> biểu đồ + bảng đạt cho cả 3 hãng; đo PDF không lỗi."""
        for h in HANG:
            with self.subTest(hang=h):
                ra = self.tmp / f'gia5_{h}.pdf'
                bc = V.ve(self.x5, h, ra)
                self.assertEqual(bc['so_trang'], 1)
                self.assertEqual(len(bc['trang'][0]['nhan']), 5)
                self.assertGreaterEqual(bc['tp_min'], 4.5)
                kho = self.HS['chung']['kho_' + self.HS['hang'][h]['kho']]
                w, hgt = bc['trang'][0]['kho']
                self.assertLess(w, kho[0])
                self.assertLess(hgt, kho[1])
                self.assertEqual(kiem_pdf(ra, h), [])
                doc = pymupdf.open(ra)
                txt = doc[0].get_text()
                gach = [sp for b in doc[0].get_text('dict')['blocks'] for ln in b.get('lines', [])
                        for sp in ln['spans'] if sp['text'].strip() == '–']
                doc.close()
                for md in self.ds5:
                    self.assertIn(md['m3'], txt)
                self.assertIn('1~ 220–240 V 50 Hz', txt)
                self.assertIn('Typical values, not guaranteed.', txt)
                self.assertEqual(len(gach), 3, 'đúng 3 ô "–" (3 model không có bản 1 pha)')

    def test_2_tu_choi_du_lieu_sai(self):
        """H tăng khi Q tăng -> dừng, báo model + điểm; CLI trả mã 2; không tạo PDF."""
        ra = self.tmp / 'sai.pdf'
        with self.assertRaises(V.LoiDuLieu) as cm:
            V.ve(self.x_sai, 'PEROTAC', ra)
        tb = str(cm.exception)
        self.assertIn('TP 50-2', tb)
        self.assertIn('Q 0 m³/h; H 20 m', tb)
        self.assertIn('Q 4 m³/h; H 21.5 m', tb)
        self.assertFalse(ra.exists())
        self.assertEqual(V.main(['--du-lieu', str(self.x_sai), '--hang', 'REDPUMP', '--ra', str(ra)]), 2)
        self.assertFalse(ra.exists())

    def test_3_tach_khi_qua_nhieu_duong(self):
        """10 model -> tách trang: PEROTAC/REDPUMP ≤ 8 đường, AQUARIS ≤ 7 (số màu)."""
        toi_da = self.HS['chung']['so_duong_toi_da_moi_bieu_do']
        for h in HANG:
            with self.subTest(hang=h):
                ra = self.tmp / f'gia10_{h}.pdf'
                bc = V.ve(self.x10, h, ra)
                self.assertGreaterEqual(bc['so_trang'], 2)
                gioi_han = min(toi_da, len(self.HS['chung']['okabe_ito_bo_vang'])) if h == 'AQUARIS' else toi_da
                for t in bc['trang']:
                    self.assertLessEqual(len(t['model']), gioi_han)
                self.assertEqual(sum(len(t['model']) for t in bc['trang']), 10)
                self.assertEqual(kiem_pdf(ra, h), [])

    def test_4_kieu_rieng_va_chon_model(self):
        """--kieu bieu_do / bang, --model chọn 2 model, khổ tự đặt."""
        for h in HANG:
            with self.subTest(hang=h):
                ra = self.tmp / f'bd_{h}.pdf'
                bc = V.ve(self.x5, h, ra, models=['TP 50-1', 'tp 55-3'], kieu='bieu_do', rong=500, cao=360)
                self.assertEqual(bc['trang'][0]['model'], ['TP 50-1', 'TP 55-3'])
                self.assertEqual(bc['trang'][0]['kho'], (500.0, 360.0))
                self.assertEqual(kiem_pdf(ra, h), [])
                ra2 = self.tmp / f'bang_{h}.pdf'
                bc2 = V.ve(self.x5, h, ra2, kieu='bang')
                self.assertNotIn('nhan', bc2['trang'][0])
                self.assertEqual(kiem_pdf(ra2, h), [])
        with self.assertRaises(V.LoiDuLieu):
            V.ve(self.x5, 'PEROTAC', self.tmp / 'x.pdf', models=['KHONG CO'])

    def test_5_tuong_phan_va_noi_suy(self):
        """Chữ trắng trên đầu bảng tối REDPUMP ≥ 4,5:1; PCHIP qua đúng điểm, không móc."""
        chinh = V.mau(self.HS['hang']['REDPUMP']['mau']['chinh'])
        self.assertGreaterEqual(V.tuong_phan(V.TRANG, chinh), 4.5)
        ra = self.tmp / 'gia5_REDPUMP_tp.pdf'
        V.ve(self.x5, 'REDPUMP', ra)
        doc = pymupdf.open(ra)
        trang = [sp for b in doc[0].get_text('dict')['blocks'] for ln in b.get('lines', []) for sp in ln['spans']
                 if sp['color'] == 0xFFFFFF and sp['text'].strip()]
        doc.close()
        self.assertTrue(any(sp['text'].strip() == 'P2' for sp in trang), 'đầu bảng REDPUMP chữ trắng')
        md = dict(khoa='X', diem=[(0, 30.0, 2), (5, 30.0, 3), (10, 26.0, 4), (15, 18.0, 5)])
        f = V.kiem_duong(md, True)
        q = np.array([0, 5, 10, 15.0])
        self.assertLess(np.max(np.abs(f(q) - np.array([30, 30, 26, 18.0]))), 1e-9)
        qq = np.linspace(0, 15, 500)
        self.assertTrue(np.all(np.diff(f(qq)) <= 1e-9))

    def test_7_phong_vendor_va_thieu_phong(self):
        """Phông lấy từ fonts/ hoặc vendor/fonts; thiếu phông hồ sơ -> mã 3; cho phép gần giống -> cảnh báo."""
        for ho in ('Lato', 'Liberation Sans'):
            ph = V.Phong(ho)
            self.assertFalse(ph.gan_giong)
            for dam in (False, True):
                self.assertIn(Path(ph.tep[dam]).parent.resolve(), [d.resolve() for d in V.thu_muc_phong_rieng()],
                              f'{ho}: phải lấy từ fonts/ hoặc vendor/fonts/, đang lấy {ph.tep[dam]}')
        hs = json.loads(json.dumps(self.HS))
        hs['hang']['AQUARIS']['phong'] = ['Khong Co Phong Nay']
        hs_gia = self.tmp / 'ho_so_thieu_phong.json'
        hs_gia.write_text(json.dumps(hs, ensure_ascii=False), encoding='utf-8')
        ra = self.tmp / 'thieu_phong.pdf'
        with self.assertRaises(V.LoiBoCuc) as cm:
            V.ve(self.x5, 'AQUARIS', ra, ho_so=hs_gia)
        self.assertIn('thiếu phông Khong Co Phong Nay', str(cm.exception))
        with contextlib.redirect_stdout(io.StringIO()):
            ma = V.main(['--du-lieu', str(self.x5), '--hang', 'AQUARIS', '--ra', str(ra), '--ho-so', str(hs_gia)])
        self.assertEqual(ma, 3)
        self.assertFalse(ra.exists())
        bc = V.ve(self.x5, 'AQUARIS', ra, ho_so=hs_gia, cho_phep_gan_giong=True)
        self.assertTrue(bc['phong_gan_giong'])
        self.assertTrue(any(c.startswith('gần giống:') for c in bc['canh_bao']))
        self.assertIn('gần giống:', Path(bc['tep_canh_bao']).read_text(encoding='utf-8'))

    def test_8_du_lieu_tao_gon(self):
        """Ca thật CM 32-160/2.2: H(Q 9) 28.5 -> 29.5 tạo đoạn phẳng rồi gãy -> mã 2; 28.5 thì đạt."""
        x_dung = tao_xlsx(self.tmp / 'cm32_dung.xlsx', ca_cm32_160_22(28.5), 'PEROTAC')
        x_gon = tao_xlsx(self.tmp / 'cm32_gon.xlsx', ca_cm32_160_22(29.5), 'PEROTAC')
        V.ve(x_dung, 'PEROTAC', self.tmp / 'cm32_dung.pdf')
        ra = self.tmp / 'cm32_gon.pdf'
        with self.assertRaises(V.LoiDuLieu) as cm:
            V.ve(x_gon, 'PEROTAC', ra)
        tb = str(cm.exception)
        for c in ('GỢN', 'CM 32-160/2.2', 'Q 9 m³/h; H 29.5 m', 'chỉnh điểm kề', 'cũ→mới', 'hỏi anh', '--chap-nhan-gon'):
            self.assertIn(c, tb)
        self.assertNotIn('CM 32-160/3', tb)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(V.main(['--du-lieu', str(x_gon), '--hang', 'PEROTAC', '--ra', str(ra)]), 2)
        bc = V.ve(x_gon, 'PEROTAC', ra, chap_nhan_gon=True)
        self.assertTrue(any('GỢN' in c and 'CM 32-160/2.2' in c for c in bc['canh_bao']))
        V.ve(x_gon, 'PEROTAC', self.tmp / 'cm32_bang.pdf', kieu='bang')        # bảng không vẽ đường: không chặn
        song = [md_tu_diem('SONG 1', [(0, 30), (2, 29.8), (4, 28), (6, 27.6), (8, 25), (10, 24.4), (12, 20)])]
        with self.assertRaises(V.LoiDuLieu) as cm2:
            V.ve(tao_xlsx(self.tmp / 'song.xlsx', song), 'REDPUMP', self.tmp / 'song.pdf')
        self.assertIn('độ cong đổi chiều', str(cm2.exception))
        for md in self.ds5:                                                   # dữ liệu trơn không bị gắn cờ
            self.assertEqual(V.kiem_gon(dict(khoa=md['m3'], diem=[(q, h, 0) for q, h in md['diem']])), [])

    def test_9_redpump_nhan_cuoi_phai(self):
        """REDPUMP: 7 đường cùng điểm cuối sát nhau -> mọi ô tên ở bên phải điểm cuối đường, sát cuối."""
        x = tao_xlsx(self.tmp / '4sp3.xlsx', ca_4sp3())
        for ra, du_lieu in ((self.tmp / '4sp3_RED.pdf', x), (self.tmp / 'gia5_RED_cuoi.pdf', self.x5)):
            bc = V.ve(du_lieu, 'REDPUMP', ra)
            self.assertEqual(kiem_pdf(ra, 'REDPUMP'), [])
            for t in bc['trang']:
                for ten, r in t['nhan']:
                    xe, ye = t['cuoi'][ten]
                    r = pymupdf.Rect(r)
                    self.assertGreaterEqual(r.x0, xe - 0.5, f'{ten}: ô không ở bên phải điểm cuối ({r} / {xe})')
                    self.assertLessEqual(r.x0 - xe, 14, f'{ten}: ô xa điểm cuối quá')
                    self.assertLessEqual(abs((r.y0 + r.y1) / 2 - ye), 4 * r.height, f'{ten}: ô lệch dọc quá xa')

    def test_10_cot_dong_dien_a(self):
        """Nguồn không có A: bỏ cột A (QĐ 152), không "hỏi anh". Có A nhưng 1 model thiếu: giữ cột + cảnh báo + tệp."""
        ds = [dict(md, a1=None, a3=None) for md in self.ds5]
        x = tao_xlsx(self.tmp / 'khong_a.xlsx', ds)
        ds2 = [dict(md, a1=None, a3=None) if i == 0 else md for i, md in enumerate(self.ds5)]
        x2 = tao_xlsx(self.tmp / 'thieu_a.xlsx', ds2)
        for h in HANG:
            with self.subTest(hang=h):
                ra = self.tmp / f'khong_a_{h}.pdf'
                bc = V.ve(x, h, ra)
                self.assertFalse(any(c.startswith('I A') for c in bc['trang'][0]['bang']['cot']))
                self.assertTrue(any('bỏ cột A (QĐ 152)' in c for c in bc['canh_bao']))
                self.assertFalse(any('hỏi anh' in c or 'THIẾU DÒNG ĐIỆN A' in c for c in bc['canh_bao']))
                self.assertEqual(kiem_pdf(ra, h), [])
                bc3 = V.ve(x, h, self.tmp / f'bo_a_{h}.pdf', bo_cot_a=True)      # cờ cũ vẫn chạy
                self.assertFalse(any(c.startswith('I A') for c in bc3['trang'][0]['bang']['cot']))
                ra2 = self.tmp / f'thieu_a_{h}.pdf'
                bc2 = V.ve(x2, h, ra2)
                self.assertTrue(any(c.startswith('I A') for c in bc2['trang'][0]['bang']['cot']))
                cb = [c for c in bc2['canh_bao'] if c.startswith('THIẾU DÒNG ĐIỆN A ở 1 model')]
                self.assertEqual(len(cb), 1)
                self.assertNotIn('hỏi anh', cb[0])
                tep = Path(str(ra2) + '.canh_bao.txt')
                self.assertTrue(tep.is_file())
                self.assertIn('THIẾU DÒNG ĐIỆN A', tep.read_text(encoding='utf-8'))
                self.assertEqual(kiem_pdf(ra2, h), [])

    def test_12_gon_so_voi_so_lieu_goc(self):
        """Gợn có sẵn trong bản gốc -> cảnh báo, không chặn; điểm đã đổi -> mã 2; tệp trong DU_LIEU/ -> so với chính nó."""
        goc_cu = V.THU_MUC_DU_LIEU
        dl = self.tmp / 'DU_LIEU'
        dl.mkdir(exist_ok=True)
        ngoai = self.tmp / 'ngoai'
        ngoai.mkdir(exist_ok=True)
        try:
            V.THU_MUC_DU_LIEU = dl
            goc_gon = tao_xlsx(dl / 'CMG.xlsx', ca_cm32_160_22(29.5), 'PEROTAC')          # bản gốc đã có gợn
            goc_tron = tao_xlsx(self.tmp / 'goc_tron.xlsx', ca_cm32_160_22(28.5), 'PEROTAC')
            vao = tao_xlsx(ngoai / 'CMG.xlsx', ca_cm32_160_22(29.5), 'PEROTAC')            # bản sửa, ngoài DU_LIEU/
            # (1) giống hệt bản gốc (--so-voi) -> chỉ cảnh báo, có trong .canh_bao.txt
            ra = self.tmp / 'gon_co_san.pdf'
            bc = V.ve(vao, 'PEROTAC', ra, so_voi=goc_gon)
            cb = [c for c in bc['canh_bao'] if 'gợn có sẵn trong số liệu đã duyệt — không chặn' in c]
            self.assertTrue(cb and 'CM 32-160/2.2' in cb[0], bc['canh_bao'])
            self.assertIn('gợn có sẵn trong số liệu đã duyệt', Path(bc['tep_canh_bao']).read_text(encoding='utf-8'))
            # không ghi --so-voi: tự lấy DU_LIEU/<cùng tên> -> cũng chỉ cảnh báo
            bc = V.ve(vao, 'PEROTAC', self.tmp / 'gon_tu_goc.pdf')
            self.assertTrue(any('gợn có sẵn' in c for c in bc['canh_bao']))
            # (2) điểm đã đổi so với bản gốc (28.5 -> 29.5) mà tạo gợn -> mã 2
            with self.assertRaises(V.LoiDuLieu) as cm:
                V.ve(vao, 'PEROTAC', self.tmp / 'gon_moi.pdf', so_voi=goc_tron)
            self.assertIn('điểm đã đổi so với số liệu gốc', str(cm.exception))
            self.assertIn('CM 32-160/2.2', str(cm.exception))
            with contextlib.redirect_stdout(io.StringIO()):
                ma = V.main(['--du-lieu', str(vao), '--hang', 'PEROTAC', '--ra', str(self.tmp / 'gon_moi.pdf'),
                             '--so-voi', str(goc_tron)])
            self.assertEqual(ma, 2)
            bc = V.ve(vao, 'PEROTAC', self.tmp / 'gon_moi2.pdf', so_voi=goc_tron, chap_nhan_gon=True)
            self.assertTrue(any(c.startswith('GỢN (vẽ bất chấp') for c in bc['canh_bao']))
            # (3) tệp vào nằm trong DU_LIEU/ -> so với chính nó -> chỉ cảnh báo
            bc = V.ve(goc_gon, 'PEROTAC', self.tmp / 'gon_trong_dl.pdf')
            self.assertTrue(any('gợn có sẵn' in c and 'CMG.xlsx' in c for c in bc['canh_bao']))
        finally:
            V.THU_MUC_DU_LIEU = goc_cu
        that = Path(V.THU_MUC_DU_LIEU) / 'CM.xlsx'                 # số liệu thật (anh đã duyệt) không bị chặn
        if that.is_file():
            ms = ['CM 80-125/4', 'CM 80-125/5.5', 'CM 80-125/7.5', 'CM 80-160/11', 'CM 80-160/15', 'CM 80-160/18.5',
                  'CM 80-160/22']
            bc = V.ve(that, 'PEROTAC', self.tmp / 'cm_that.pdf', models=ms)
            self.assertTrue(any('gợn có sẵn' in c for c in bc['canh_bao']))

    def test_11_help_aquaris_7_duong(self):
        """--help ghi rõ AQUARIS tối đa 7 đường/biểu đồ."""
        out = io.StringIO()
        with contextlib.redirect_stdout(out), self.assertRaises(SystemExit):
            V.main(['--help'])
        self.assertIn('AQUARIS 7', out.getvalue())
        self.assertIn('--so-voi', out.getvalue())
        self.assertIn('gợn có sẵn trong số liệu đã duyệt', out.getvalue())
        self.assertIn('tối đa 7', out.getvalue())

    def test_6_bo_kiem_bat_duoc_loi(self):
        """Đối chứng: PDF cố ý có chữ chồng chữ và nét đường cong đè chữ -> kiem_pdf phải báo."""
        ph = V.Phong('Liberation Sans')
        doc = pymupdf.open()
        pg = doc.new_page(width=300, height=200)
        pg.insert_font(fontname='F', fontfile=ph.tep[False])
        pg.insert_text((50, 50), 'CM 80-125/4', fontname='F', fontsize=8)
        pg.insert_text((60, 52), 'CM 80-160/11', fontname='F', fontsize=8)
        pg.insert_text((150, 120), 'CM 80-160/22', fontname='F', fontsize=8)
        pg.draw_polyline([(140, 110), (200, 125)], color=(0.1, 0.3, 0.6), width=1.75)
        ra = self.tmp / 'doi_chung.pdf'
        doc.save(ra)
        doc.close()
        loi = kiem_pdf(ra, 'PEROTAC')
        self.assertTrue(any('chữ chồng' in x and 'CM 80-125/4' in x for x in loi), loi)
        self.assertTrue(any('đè nét đường cong' in x and 'CM 80-160/22' in x for x in loi), loi)

    def test_13_spline_mem_chi_khi_khong_doi_dang(self):
        """QĐ 158: đường mềm (spline) khi chỉ làm mềm góc; spline đổi dáng (AQUARIS 6SP46-8, DU_LIEU/SP.xlsx) -> giữ PCHIP."""
        ss = [(0, 6), (1, 5.9), (2, 5.5), (3, 4.9), (4, 4.0), (5, 2.9), (6, 1.6), (7, 0)]
        f = V.kiem_duong(dict(khoa='SS-150', diem=[(q, h, i) for i, (q, h) in enumerate(ss)]), True)
        self.assertEqual(type(f).__name__, 'CubicSpline')
        qq = np.linspace(0, 7, 500)
        self.assertTrue(np.all(np.diff(f(qq)) <= 1e-9))
        self.assertLessEqual(f(qq).max(), 6 + 1e-9)
        sp = [(0, 106), (35, 80), (40, 77), (46, 70), (50, 64), (55, 56)]
        f = V.kiem_duong(dict(khoa='6SP46-8', diem=[(q, h, i) for i, (q, h) in enumerate(sp)]), True)
        self.assertEqual(type(f).__name__, 'PchipInterpolator')


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except (AttributeError, ValueError):
        pass
    argv = [a for a in sys.argv if a != '--giu']
    kq = unittest.main(argv=argv, exit=False, verbosity=2).result
    tong = kq.testsRun
    hong = len(kq.failures) + len(kq.errors)
    print(f'\nKẾT QUẢ: {tong - hong}/{tong} ca kiểm ĐẠT' + ('' if not hong else f' — {hong} KHÔNG ĐẠT'))
    sys.exit(1 if hong else 0)
